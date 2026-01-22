import openpyxl
wb = openpyxl.load_workbook('result/keywordList_all.xlsx')
ws = wb['recent30days_sorted']
# 모든 F열을 None으로 설정해서 테스트
for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=6).value = None
wb.save('result/keywordList_all.xlsx')
print('모든 F열을 None으로 설정했습니다.')