from openpyxl import load_workbook

wb = load_workbook('result/keywordList_all.xlsx')

# recent30days_sorted 시트 확인
if 'recent30days_sorted' in wb.sheetnames:
    ws = wb['recent30days_sorted']
    print('=== recent30days_sorted 시트 확인 ===')
else:
    ws = wb['recent30days']
    print('=== recent30days 시트 확인 (recent30days_sorted 없음) ===')
print('총 행 수: {}'.format(ws.max_row))

print('\n헤더:')
for col in range(1, ws.max_column + 1):
    header = ws.cell(row=1, column=col).value
    col_letter = chr(64+col)
    print('  {}열: {}'.format(col_letter, header))

print('\n처음 10행 데이터:')
for row in range(2, min(12, ws.max_row + 1)):
    c_val = ws.cell(row=row, column=3).value
    d_val = ws.cell(row=row, column=4).value
    e_val = ws.cell(row=row, column=5).value
    f_val = ws.cell(row=row, column=6).value
    print('행{}: C={}, D={}, E={}, F={}'.format(row, c_val, d_val, e_val, f_val))

print('\n노란색 강조된 행들:')
yellow_count = 0
for row in range(2, ws.max_row + 1):
    cell_fill = ws.cell(row=row, column=1).fill
    if cell_fill and hasattr(cell_fill, 'start_color') and cell_fill.start_color and cell_fill.start_color.rgb == 'FFFFFF00':
        c_val = ws.cell(row=row, column=3).value
        d_val = ws.cell(row=row, column=4).value
        e_val = ws.cell(row=row, column=5).value
        f_val = ws.cell(row=row, column=6).value
        print('행{}: C={}, D={}, E={}, F={}'.format(row, c_val, d_val, e_val, f_val))
        yellow_count += 1

print('\n총 노란색 강조 행 수: {}'.format(yellow_count))

# 새로운 조건 검증
print('\n조건 검증 (C > 500 and D < 100 and E < 10):')
valid_count = 0
for row in range(2, ws.max_row + 1):
    c_val = ws.cell(row=row, column=3).value
    d_val = ws.cell(row=row, column=4).value
    e_val = ws.cell(row=row, column=5).value

    try:
        c_num = float(c_val) if c_val is not None else 0
        d_num = float(d_val) if d_val is not None and d_val != '100+' else 100
        e_num = float(e_val) if e_val is not None else 0

        if c_num > 500 and d_num < 100 and e_num < 10:
            valid_count += 1
            if valid_count <= 3:  # 처음 3개만 출력
                print('  행{}: C={}, D={}, E={}'.format(row, c_val, d_val, e_val))
    except (ValueError, TypeError):
        continue

print('조건에 맞는 총 행 수: {}'.format(valid_count))
print('노란색 강조 행 수와 조건 일치: {}'.format('YES' if yellow_count == valid_count else 'NO'))