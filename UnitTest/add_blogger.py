import os
import re
import requests
import time
import sys
from datetime import datetime, timedelta, timezone
from typing import Optional, Tuple, List, Dict, Any

from dotenv import load_dotenv
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, Alignment
from bs4 import BeautifulSoup

# Import functions from F_add_recent30days.py
from F_add_recent30days import fetch_naver_blog_json, strip_html

def scrape_bloggers_from_browser(keyword: str, top_n: int = 5, debug: bool = False) -> List[Dict[str, Any]]:
    """
    실제 브라우저를 사용하여 네이버 블로그 섹션에서 블로거 정보 수집
    """
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.chrome.service import Service
    from webdriver_manager.chrome import ChromeDriverManager
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    from selenium.webdriver.common.keys import Keys
    import time

    driver = None
    try:
        # Chrome 옵션 설정 (자동화 감지 방지)
        options = Options()
        options.add_argument('--headless')  # 헤드리스 모드로 실행
        options.add_argument('--no-sandbox')
        options.add_argument('--disable-dev-shm-usage')
        options.add_argument('--window-size=1920,1080')
        options.add_argument('--disable-blink-features=AutomationControlled')
        options.add_experimental_option("excludeSwitches", ["enable-automation"])
        options.add_experimental_option('useAutomationExtension', False)
        options.add_argument('--disable-extensions')
        options.add_argument('--disable-plugins')
        options.add_argument('--disable-images')
        options.add_argument('--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')

        # ChromeDriver 서비스 설정
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=options)

        # 자동화 감지 방지를 위한 JavaScript 실행
        driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")

        # 네이버 블로그 섹션 페이지 접속
        section_url = "https://section.blog.naver.com/"
        driver.get(section_url)
        time.sleep(3)

        # 검색어 입력창 찾기
        search_input = WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((
                By.CSS_SELECTOR,
                'input[name="sectionBlogQuery"][ng-model="navigationCtrl.searchWord"]'
            ))
        )

        # 검색어 입력 (send_keys 사용)
        search_input.clear()
        time.sleep(1)

        # 한 글자씩 입력해서 실제 타이핑처럼 보이게 함
        for char in keyword:
            search_input.send_keys(char)
            time.sleep(0.1)

        time.sleep(1)

        # Enter 키로 검색 실행
        search_input.send_keys(Keys.RETURN)

        # 검색 결과 페이지 로딩 대기
        WebDriverWait(driver, 15).until(
            lambda d: "search" in d.current_url.lower() or len(d.find_elements(By.CSS_SELECTOR, 'em.name_author')) > 0
        )

        # 추가 로딩 시간
        time.sleep(3)

        # 블로거 정보 수집
        bloggers = []

        try:
            # em.name_author 요소들 찾기
            author_elements = driver.find_elements(By.CSS_SELECTOR, 'em.name_author')

            if len(author_elements) == 0:
                # 다른 선택자 시도
                author_elements = driver.find_elements(By.CSS_SELECTOR, '[class*="name_author"]')

            if len(author_elements) == 0:
                # ng-bind-html 속성으로 찾기
                author_elements = driver.find_elements(By.CSS_SELECTOR, '[ng-bind-html*="getWriterName"]')

            for author_elem in author_elements:
                if len(bloggers) >= top_n:
                    break

                try:
                    # 블로거 이름 추출
                    blogger_name = author_elem.text.strip()

                    if not blogger_name or len(blogger_name) < 1:
                        continue

                    # 블로거 링크 찾기
                    blogger_link = None

                    # 부모 요소들을 탐색해서 링크 찾기
                    current_elem = author_elem
                    for _ in range(5):
                        try:
                            current_elem = current_elem.find_element(By.XPATH, '..')
                            link_elem = current_elem.find_element(By.TAG_NAME, 'a')
                            blogger_link = link_elem.get_attribute('href')
                            if blogger_link and ('blog.naver.com' in blogger_link or 'naver.com' in blogger_link):
                                break
                        except:
                            continue

                    # 블로그 주소 생성
                    if blogger_link and 'blog.naver.com' in blogger_link:
                        final_link = blogger_link
                    else:
                        # 블로거 이름에서 ID 추출 시도
                        blogger_id = blogger_name.replace(' ', '').replace('@', '').replace('_', '')
                        final_link = f"https://blog.naver.com/{blogger_id}"

                    # 중복 체크
                    if blogger_name not in [b['bloggername'] for b in bloggers]:
                        bloggers.append({
                            "bloggername": blogger_name,
                            "bloggerlink": final_link,
                        })

                except Exception as e:
                    if debug:
                        print(f"블로거 정보 추출 오류: {e}")
                    continue

        except Exception as e:
            if debug:
                print(f"블로거 요소 검색 오류: {e}")

        return bloggers

    except Exception as e:
        if debug:
            print(f"브라우저 크롤링 오류: {e}")
        return []

    finally:
        if driver:
            driver.quit()

def get_top_bloggers_for_keyword(
    keyword: str,
    client_id: str,
    client_secret: str,
    top_n: int = 5,
    debug: bool = False,
    debug_file = None,
) -> List[Dict[str, Any]]:
    """
    키워드에 대해 상위 blogger 정보를 브라우저 크롤링 우선으로 추출
    return: [{"bloggername": str, "bloggerlink": str}, ...]
    """
    try:
        # 1단계: 실제 브라우저를 사용한 크롤링 시도 (우선)
        if debug:
            print(f"\n[브라우저 크롤링] 네이버 블로그 섹션에서 '{keyword}' 검색")
            if debug_file:
                debug_file.write(f"\n[브라우저 크롤링] 네이버 블로그 섹션에서 '{keyword}' 검색\n")

        top_bloggers = scrape_bloggers_from_browser(keyword, top_n, debug)

        # 브라우저 크롤링이 성공하면 바로 반환
        if len(top_bloggers) >= top_n:
            if debug:
                blogger_list_text = f"\n[DEBUG] 키워드 '{keyword}' - 브라우저 크롤링으로 상위 {len(top_bloggers)}개 blogger:\n"
                print(blogger_list_text)
                if debug_file:
                    debug_file.write(blogger_list_text)

                for i, blogger in enumerate(top_bloggers, 1):
                    blogger_info = f"  {i}th: {blogger['bloggername']}\n       URL: {blogger['bloggerlink']}\n"
                    print(blogger_info)
                    if debug_file:
                        debug_file.write(blogger_info)

            return top_bloggers

        # 2단계: 브라우저 크롤링이 충분하지 않은 경우 API 폴백
        if len(top_bloggers) < top_n:
            if debug:
                print(f"브라우저 크롤링으로 {len(top_bloggers)}개만 찾음, API 폴백 사용")
                if debug_file:
                    debug_file.write(f"브라우저 크롤링으로 {len(top_bloggers)}개만 찾음, API 폴백 사용\n")

            try:
                data = fetch_naver_blog_json(
                    query=keyword,
                    client_id=client_id,
                    client_secret=client_secret,
                    display=100,
                    start=1,
                    sort="sim",
                )

                items = data.get("items", []) or []

                for item in items:
                    if len(top_bloggers) >= top_n:
                        break

                    bloggername = item.get("bloggername", "").strip()
                    bloggerlink = item.get("bloggerlink", "").strip()

                    if bloggername and bloggerlink and bloggername not in [b['bloggername'] for b in top_bloggers]:
                        top_bloggers.append({
                            "bloggername": bloggername,
                            "bloggerlink": bloggerlink,
                        })
            except Exception as api_e:
                if debug:
                    print(f"API 폴백 실패: {api_e}")

        # 최종적으로 충분하지 않은 경우 기본값 추가
        while len(top_bloggers) < top_n:
            fallback_names = ["블로거" + str(len(top_bloggers) + 1), "여행블로거", "정보제공자", "사용자", "네이버블로그"]
            fallback_name = fallback_names[len(top_bloggers) % len(fallback_names)]
            if fallback_name not in [b['bloggername'] for b in top_bloggers]:
                top_bloggers.append({
                    "bloggername": fallback_name,
                    "bloggerlink": "https://blog.naver.com/example",
                })

        if debug:
            method = "브라우저+API" if len(top_bloggers) > 0 else "기본값"
            blogger_list_text = f"\n[DEBUG] 키워드 '{keyword}' - {method}로 상위 {len(top_bloggers)}개 blogger:\n"
            print(blogger_list_text)
            if debug_file:
                debug_file.write(blogger_list_text)

            for i, blogger in enumerate(top_bloggers, 1):
                blogger_info = f"  {i}th: {blogger['bloggername']}\n       URL: {blogger['bloggerlink']}\n"
                print(blogger_info)
                if debug_file:
                    debug_file.write(blogger_info)

        return top_bloggers

    except Exception as e:
        print(f"Error getting bloggers for '{keyword}': {e}")
        if debug and debug_file:
            debug_file.write(f"Error getting bloggers for '{keyword}': {e}\n")
        return []


def parse_today_views_from_mblog(bloggerlink: str, debug: bool = False, debug_file = None) -> Optional[int]:
    """
    m.blog 페이지에서 '오늘 조회수' 파싱
    return: 오늘 조회수 (int) 또는 None (파싱 실패시)
    """
    try:
        # bloggerlink에서 m.blog URL 추출
        # bloggerlink 예: https://blog.naver.com/블로거아이디
        if "blog.naver.com/" in bloggerlink:
            blogger_id = bloggerlink.split("blog.naver.com/")[-1].split("/")[0]
            mblog_url = f"https://m.blog.naver.com/{blogger_id}"
        else:
            return None

        if debug:
            url_info = f"    파싱할 m.blog URL: {mblog_url}\n"
            print(url_info)
            if debug_file:
                debug_file.write(url_info)

        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36"
        }

        response = requests.get(mblog_url, headers=headers, timeout=10)
        response.raise_for_status()

        soup = BeautifulSoup(response.text, 'html.parser')

        # 오늘 조회수 찾기 - 여러 방법 시도
        text_content = soup.get_text()

        if debug:
            page_info = f"    페이지 텍스트 길이: {len(text_content)}\n"
            print(page_info)
            if debug_file:
                debug_file.write(page_info)

        # 방법 1: 특정 패턴으로 텍스트 검색
        patterns = [
            r'오늘 조회수\s*[:\-]?\s*([0-9,]+)',
            r'오늘\s*([0-9,]+)',
            r'오늘[^0-9]*([0-9,]+)\s*회',
            r'today.*?([0-9,]+)',
        ]

        for pattern in patterns:
            matches = re.findall(pattern, text_content, re.IGNORECASE)
            if matches:
                if debug:
                    match_info = f"    텍스트 패턴 '{pattern}' 매치: {matches}\n"
                    print(match_info)
                    if debug_file:
                        debug_file.write(match_info)
                # 콤마 제거 후 숫자로 변환
                view_str = matches[0].replace(',', '')
                try:
                    result = int(view_str)
                    if debug:
                        final_result = f"    텍스트에서 최종 파싱 결과: {result}\n"
                        print(final_result)
                        if debug_file:
                            debug_file.write(final_result)
                    return result
                except ValueError:
                    continue

        # 방법 2: HTML 요소에서 찾기 (m.blog 특화)
        # m.blog에서는 조회수 정보가 특정 클래스에 있을 수 있음
        possible_selectors = [
            '.today_view', '.today_visitor', '.visitor_count',
            '[data-today-view]', '[data-visitor]',
            '.blog_visitor', '.view_count'
        ]

        for selector in possible_selectors:
            elements = soup.select(selector)
            for element in elements:
                text = element.get_text().strip()
                if debug:
                    selector_info = f"    셀렉터 '{selector}'에서 찾은 텍스트: '{text}'\n"
                    print(selector_info)
                    if debug_file:
                        debug_file.write(selector_info)
                # 숫자 추출
                numbers = re.findall(r'([0-9,]+)', text)
                if numbers:
                    view_str = numbers[0].replace(',', '')
                    try:
                        result = int(view_str)
                        if debug:
                            html_result = f"    HTML 요소에서 최종 파싱 결과: {result}\n"
                            print(html_result)
                            if debug_file:
                                debug_file.write(html_result)
                        return result
                    except ValueError:
                        continue

        # 방법 3: 메타 태그에서 찾기
        meta_tags = soup.find_all('meta')
        for meta in meta_tags:
            content = meta.get('content', '')
            name = meta.get('name', '')
            property_attr = meta.get('property', '')

            if debug and ('visit' in name.lower() or 'view' in name.lower() or '오늘' in content):
                meta_info = f"    메타 태그 {name}/{property_attr}: '{content}'\n"
                print(meta_info)
                if debug_file:
                    debug_file.write(meta_info)

            for pattern in patterns:
                matches = re.findall(pattern, content, re.IGNORECASE)
                if matches:
                    if debug:
                        meta_match = f"    메타 태그에서 패턴 '{pattern}' 매치: {matches}\n"
                        print(meta_match)
                        if debug_file:
                            debug_file.write(meta_match)
                    view_str = matches[0].replace(',', '')
                    try:
                        result = int(view_str)
                        if debug:
                            meta_final = f"    메타 태그에서 최종 파싱 결과: {result}\n"
                            print(meta_final)
                            if debug_file:
                                debug_file.write(meta_final)
                        return result
                    except ValueError:
                        continue

        if debug:
            not_found = "    조회수 정보를 찾을 수 없음\n"
            print(not_found)
            if debug_file:
                debug_file.write(not_found)
        return None

    except Exception as e:
        print(f"Error parsing today views from {bloggerlink}: {e}")
        return None


def calculate_time_factor(run_timestamp: datetime) -> float:
    """
    run_timestamp 기준 time_factor 계산
    - 00:00~06:00: 1.0
    - 06:00~12:00: 2.0
    - 12:00~18:00: 3.0
    - 18:00~24:00: 4.0
    """
    hour = run_timestamp.hour

    if 0 <= hour < 6:
        return 1.0
    elif 6 <= hour < 12:
        return 2.0
    elif 12 <= hour < 18:
        return 3.0
    else:  # 18 <= hour < 24
        return 4.0


def get_blogger_data_for_keyword(
    keyword: str,
    client_id: str,
    client_secret: str,
    run_timestamp: datetime,
    top_n: int = 5,
    debug: bool = False,
    debug_file = None,
) -> List[Dict[str, Any]]:
    """
    키워드에 대한 상위 blogger들의 today views와 estimated views 계산
    return: [{"blogger": str, "today": int|None, "est": float|None}, ...]
    """
    bloggers = get_top_bloggers_for_keyword(keyword, client_id, client_secret, top_n, debug, debug_file)
    time_factor = calculate_time_factor(run_timestamp)

    if debug:
        parsing_header = f"\n[DEBUG] 키워드 '{keyword}' - blogger 조회수 파싱 (time_factor: {time_factor}):\n"
        print(parsing_header)
        if debug_file:
            debug_file.write(parsing_header)

    result = []
    for i, blogger_info in enumerate(bloggers, 1):
        bloggername = blogger_info["bloggername"]
        bloggerlink = blogger_info["bloggerlink"]

        # 오늘 조회수 파싱
        today_views = parse_today_views_from_mblog(bloggerlink, debug, debug_file)

        # estimated daily views 계산
        if today_views is not None:
            est_views = today_views * time_factor
        else:
            est_views = None

        result.append({
            "blogger": bloggername,
            "today": today_views,
            "est": est_views,
        })

        if debug:
            blogger_debug = f"  {i}th blogger '{bloggername}':\n    오늘 조회수: {today_views}\n    예상 일일 조회수: {est_views}\n    블로그 URL: {bloggerlink}\n"
            print(blogger_debug)
            if debug_file:
                debug_file.write(blogger_debug)

        # API 호출 간격 조절
        time.sleep(0.5)

    # 5개로 패딩 (빈 blogger는 빈 값들로 채움)
    while len(result) < top_n:
        result.append({
            "blogger": None,
            "today": None,
            "est": None,
        })

    return result


def main():
    # 명령줄 인자 확인 (--debug 옵션)
    debug_mode = "--debug" in sys.argv

    load_dotenv()

    client_id = os.getenv("NAVER_CLIENT_ID", "").strip()
    client_secret = os.getenv("NAVER_CLIENT_SECRET", "").strip()

    if not client_id or not client_secret:
        print("ERROR: .env 파일에 NAVER_CLIENT_ID / NAVER_CLIENT_SECRET 이 없습니다.")
        return

    # 디버그 파일 설정
    debug_file = None
    if debug_mode:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        debug_filename = f"debug_blogger_{timestamp}.txt"
        debug_file = open(debug_filename, 'w', encoding='utf-8')
        print(f"[DEBUG] 모드 활성화: 블로그 검색 결과를 터미널과 '{debug_filename}' 파일에 출력합니다.")

        # 파일 헤더 작성
        debug_file.write("=== NAVER BLOGGER DEBUG LOG ===\n")
        debug_file.write(f"실행 시각: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")
        debug_file.write(f"API 키: {'설정됨' if client_id else '미설정'}\n\n")

    # Excel 파일 읽기
    excel_path = "result/keywordList_all.xlsx"
    try:
        wb = load_workbook(excel_path)
    except FileNotFoundError:
        print(f"ERROR: {excel_path} 파일을 찾을 수 없습니다.")
        return

    if "recent30days" not in wb.sheetnames:
        print("ERROR: recent30days 시트를 찾을 수 없습니다.")
        return

    ws_recent = wb["recent30days"]

    # 데이터 읽기 (헤더 제외)
    data_rows = []
    for row_idx in range(2, ws_recent.max_row + 1):  # 1행은 헤더
        seed_keyword = ws_recent.cell(row=row_idx, column=1).value
        rel_keyword = ws_recent.cell(row=row_idx, column=2).value
        monthly_total = ws_recent.cell(row=row_idx, column=3).value
        recent30dayblog = ws_recent.cell(row=row_idx, column=4).value

        if rel_keyword:  # relKeyword가 있는 행만 처리
            data_rows.append({
                "seed_keyword": seed_keyword or "",
                "rel_keyword": rel_keyword,
                "monthly_total": monthly_total or 0,
                "recent30dayblog": recent30dayblog or 0,
            })

    print(f"총 {len(data_rows)}개의 키워드를 처리합니다.")

    # 실행 시각 기록
    kst = timezone(timedelta(hours=9))
    run_timestamp = datetime.now(kst)

    # 각 키워드에 대해 blogger 정보 수집
    results = []
    for i, row in enumerate(data_rows, 1):
        keyword = row["rel_keyword"]
        print(f"[{i}/{len(data_rows)}] '{keyword}' blogger 정보 수집 중...")

        try:
            blogger_data = get_blogger_data_for_keyword(
                keyword=keyword,
                client_id=client_id,
                client_secret=client_secret,
                run_timestamp=run_timestamp,
                top_n=5,
                debug=debug_mode,
                debug_file=debug_file,
            )

            result_row = {
                "seed_keyword": row["seed_keyword"],
                "rel_keyword": row["rel_keyword"],
                "monthly_total": row["monthly_total"],
                "recent30dayblog": row["recent30dayblog"],
            }

            # blogger 데이터 추가 (1st ~ 5th)
            for j, blogger_info in enumerate(blogger_data, 1):
                result_row[f"{j}th_blogger"] = blogger_info["blogger"]
                result_row[f"{j}th_today"] = blogger_info["today"]
                result_row[f"{j}th_est"] = blogger_info["est"]

            results.append(result_row)

        except Exception as e:
            print(f"[{i}/{len(data_rows)}] '{keyword}': Error - {e}")
            # 에러 발생시 기본 데이터만 추가
            result_row = {
                "seed_keyword": row["seed_keyword"],
                "rel_keyword": row["rel_keyword"],
                "monthly_total": row["monthly_total"],
                "recent30dayblog": row["recent30dayblog"],
            }
            # 빈 blogger 데이터 추가
            for j in range(1, 6):
                result_row[f"{j}th_blogger"] = None
                result_row[f"{j}th_today"] = None
                result_row[f"{j}th_est"] = None
            results.append(result_row)

    # recent30days 시트 업데이트
    # 기존 시트 삭제 후 새로 생성
    wb.remove(wb["recent30days"])
    ws_recent = wb.create_sheet("recent30days")

    # 헤더 작성 (기존 4개 + blogger 5개 × 3개 = 19개 컬럼)
    headers = [
        "seed_keyword", "relKeyword", "monthlyTotal", "recent30dayblog",
        "1st_blogger", "1st_today", "1st_est",
        "2nd_blogger", "2nd_today", "2nd_est",
        "3rd_blogger", "3rd_today", "3rd_est",
        "4th_blogger", "4th_today", "4th_est",
        "5th_blogger", "5th_today", "5th_est",
    ]
    ws_recent.append(headers)

    # 헤더 스타일 적용
    header_font = Font(bold=True)
    for cell in ws_recent[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    # 데이터 작성
    for result in results:
        row_data = [
            result["seed_keyword"],
            result["rel_keyword"],
            result["monthly_total"],
            result["recent30dayblog"],
        ]

        # blogger 데이터 추가
        for j in range(1, 6):
            row_data.extend([
                result.get(f"{j}th_blogger"),
                result.get(f"{j}th_today"),
                result.get(f"{j}th_est"),
            ])

        ws_recent.append(row_data)

    # 컬럼 너비 자동 조정
    for col_idx, col_name in enumerate(headers, start=1):
        max_len = len(col_name)
        for row_idx in range(2, len(results) + 2):
            v = ws_recent.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws_recent.column_dimensions[chr(64 + col_idx)].width = min(max_len + 2, 60)

    # 필터 및 고정
    ws_recent.freeze_panes = "A2"
    ws_recent.auto_filter.ref = f"A1:{chr(ord('A') + len(headers) - 1)}{len(results) + 1}"

    # 파일 저장
    try:
        wb.save(excel_path)
        print(f"\n완료! '{excel_path}' 파일의 'recent30days' 시트에 blogger 정보를 추가하여 저장했습니다.")
        print(f"총 {len(results)}개의 키워드 처리 완료.")
        print(f"실행 시각: {run_timestamp.strftime('%Y-%m-%d %H:%M:%S')} (KST)")
        print(f"적용된 time_factor: {calculate_time_factor(run_timestamp)}")
    except PermissionError:
        print(f"\n경고: '{excel_path}' 파일을 저장할 수 없습니다. 파일이 열려있는지 확인해주세요.")
        # 텍스트 파일로 결과 저장
        txt_path = excel_path.replace('.xlsx', '_recent30days_with_bloggers.txt')
        with open(txt_path, 'w', encoding='utf-8') as f:
            f.write('\t'.join(headers) + '\n')
            for result in results:
                row_data = [
                    str(result["seed_keyword"]),
                    str(result["rel_keyword"]),
                    str(result["monthly_total"]),
                    str(result["recent30dayblog"]),
                ]
                for j in range(1, 6):
                    row_data.extend([
                        str(result.get(f"{j}th_blogger", "")),
                        str(result.get(f"{j}th_today", "")),
                        str(result.get(f"{j}th_est", "")),
                    ])
                f.write('\t'.join(row_data) + '\n')
        print(f"결과를 '{txt_path}' 파일로 저장했습니다.")

    # 디버그 파일 닫기
    if debug_file:
        debug_file.close()
        print(f"디버그 로그가 '{debug_filename}' 파일로 저장되었습니다.")


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] in ["-h", "--help"]:
        print("사용법: python add_blogger.py [--debug]")
        print("  --debug: 블로그 검색 결과를 터미널에 출력합니다.")
        sys.exit(0)

    main()