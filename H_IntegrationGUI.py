"""
H_IntegrationGUI.py - 네이버 키워드 분석 통합 GUI
A~G 모듈을 순차적으로 실행하는 통합 인터페이스

실행 순서:
A_rightside → B_autocomplete → C_sumKeyword → D_searchresult → E_deleteAndPriority → F_add_recent30days → F_sort_result_1 → G_add_blogger_by_mainPage
"""

import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext, filedialog
import threading
import queue
import sys
import os
import json
import time
import random
from typing import List, Optional
from openpyxl import load_workbook

# 황금키워드 선별 기준: 최근 30일 발행 블로그 수가 이 값 미만인 키워드를 이 개수만큼 찾는다.
GOLDEN_BLOG_THRESHOLD = 20
GOLDEN_TARGET_COUNT = 50

# GUI 설정(마지막으로 사용한 결과 폴더 등)을 저장하는 파일. 스크립트 위치 기준 고정 경로를 사용해
# 실행 중 결과 폴더로 작업 디렉터리(cwd)가 바뀌어도 항상 같은 설정 파일을 찾을 수 있도록 한다.
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_FILE = os.path.join(SCRIPT_DIR, "gui_config.json")
# 키워드 처리 대기열/현황을 저장해 프로그램을 재시작해도 이어서 처리할 수 있게 하는 파일
TASKS_FILE = os.path.join(SCRIPT_DIR, "gui_tasks.json")


def load_last_output_dir() -> str:
    """마지막으로 사용한 결과 저장 폴더를 불러온다. 없거나 유효하지 않으면 현재 폴더를 반환."""
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            path = data.get("last_output_dir", "")
            if path and os.path.isdir(path):
                return path
    except Exception:
        pass
    return os.getcwd()


def save_last_output_dir(path: str):
    """결과 저장 폴더를 다음 실행에도 재사용할 수 있도록 저장한다."""
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump({"last_output_dir": path}, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"설정 저장 실패: {e}")

# 모듈 import
import A_rightside
import B_autocomplete
import C_sumKeyword
import D_searchresult
import E_deleteAndPriority
import F_add_recent30days
import F_sort_result_1
import G_add_blogger_by_mainPage


class MockInput:
    """GUI에서 모듈 input을 모의하는 클래스"""
    def __init__(self, seed_keyword: str = "", required_keywords: List[str] = None):
        self.seed_keyword = seed_keyword  # 단일 키워드
        self.required_keywords = required_keywords or []
        self.seed_call_count = 0
        self.required_call_count = 0

    def __call__(self, prompt):
        if "seed >" in prompt:
            self.seed_call_count += 1
            if self.seed_call_count == 1:
                # 첫 번째 호출: 단일 시드 키워드 반환
                return self.seed_keyword
            else:
                # 두 번째 호출: 빈 입력으로 종료
                return ""
        elif "필수 키워드 >" in prompt:
            self.required_call_count += 1
            if self.required_call_count == 1:
                # 필수 키워드들 반환
                return ", ".join(self.required_keywords)
        return input(prompt)  # 실제 input으로 폴백


class Logger:
    """GUI 텍스트 위젯과 파일에 동시에 출력하는 로거 클래스"""
    def __init__(self, text_widget=None, log_file="log.txt"):
        self.text_widget = text_widget
        self.log_file = log_file
        self.buffer = ""

        # 로그 파일 초기화
        try:
            with open(self.log_file, 'w', encoding='utf-8') as f:
                f.write("=== 네이버 키워드 분석 로그 시작 ===\n")
        except Exception as e:
            print(f"로그 파일 초기화 실패: {e}")

    def write(self, text):
        """텍스트를 GUI와 파일에 동시에 출력"""
        # 버퍼에 추가
        self.buffer += text

        # 모든 완전한 줄 처리
        lines = self.buffer.split('\n')
        self.buffer = lines.pop() if lines and not self.buffer.endswith('\n') else ''

        # 완전한 줄들을 처리
        for line in lines:
            line += '\n'

            # GUI에 출력 (스레드 안전하게)
            if self.text_widget:
                self.text_widget.after(0, lambda l=line: self._append_to_text_widget(l))

            # 파일에 출력
            try:
                with open(self.log_file, 'a', encoding='utf-8') as f:
                    f.write(line)
            except Exception as e:
                # 파일 쓰기 실패 시 GUI에 에러 표시
                if self.text_widget:
                    error_msg = f"[로그 파일 쓰기 오류: {e}]\n"
                    self.text_widget.after(0, lambda: self._append_to_text_widget(error_msg))

    def flush(self):
        """버퍼 비우기"""
        if self.buffer:
            # 남은 버퍼 내용 처리
            if self.text_widget:
                self.text_widget.after(0, lambda: self._append_to_text_widget(self.buffer))

            # 파일에 출력
            try:
                with open(self.log_file, 'a', encoding='utf-8') as f:
                    f.write(self.buffer)
            except Exception as e:
                if self.text_widget:
                    error_msg = f"[로그 파일 쓰기 오류: {e}]\n"
                    self.text_widget.after(0, lambda: self._append_to_text_widget(error_msg))

        self.buffer = ""

    def _append_to_text_widget(self, text):
        """GUI 텍스트 위젯에 텍스트 추가 (메인 스레드에서 실행)"""
        if self.text_widget:
            self.text_widget.insert(tk.END, text)
            self.text_widget.see(tk.END)  # 자동 스크롤


class KeywordTaskQueue:
    """키워드 처리 대기열.

    queue.Queue와 달리 대기 중인 항목을 임의로 삭제하거나 순서를 바꿀 수 있어야 하므로,
    내부 리스트를 트리 위젯(표시 순서/상태의 기준)과 맞춰 통째로 교체하는 방식으로 동작한다.
    """

    def __init__(self):
        self._items: List[tuple] = []
        self._lock = threading.Lock()
        self._cv = threading.Condition(self._lock)

    def replace_all(self, items: List[tuple]):
        """대기 목록을 통째로 교체한다 (삭제/순서 변경/초기 복원 시 사용)"""
        with self._cv:
            self._items = list(items)
            if self._items:
                self._cv.notify_all()

    def get(self, timeout: Optional[float] = None):
        with self._cv:
            if not self._items:
                self._cv.wait(timeout=timeout)
            if not self._items:
                raise queue.Empty
            return self._items.pop(0)

    def empty(self) -> bool:
        with self._lock:
            return len(self._items) == 0

    def task_done(self):
        """queue.Queue와의 호환을 위한 자리표시자 (별도 동작 없음)"""
        pass


class IntegrationGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("네이버 키워드 분석 통합 시스템 v1.0")
        self.root.geometry("1280x720")
        self.root.minsize(1000, 600)
        self.root.resizable(True, True)

        # 실행 상태 관리
        self.is_running = False
        self.stop_requested = False

        # 키워드 처리 대기열 (실행 중에도 키워드 추가/삭제/순서 변경 가능)
        self.keyword_queue = KeywordTaskQueue()
        self._keyword_counter = 0
        self.processing_iid: Optional[str] = None  # 현재 처리 중인 항목 (대기열 재구성 시 제외하기 위함)

        # GUI 구성
        self.create_widgets()

        # Logger 초기화 및 sys.stdout 대체
        self.logger = Logger(text_widget=self.log_text)
        self.original_stdout = sys.stdout
        sys.stdout = self.logger

        # 이전 실행에서 저장된 키워드 처리 현황을 불러와 이어서 진행할 수 있도록 복원
        self.load_tasks()


    def create_widgets(self):
        """GUI 위젯 생성"""
        # 메인 프레임
        main_frame = ttk.Frame(self.root, padding="6")
        main_frame.pack(fill=tk.BOTH, expand=True)

        # 타이틀
        title_label = ttk.Label(main_frame, text="네이버 키워드 분석 통합 시스템",
                               font=("Arial", 12, "bold"))
        title_label.pack(pady=(0, 6))

        # 입력 프레임
        input_frame = ttk.LabelFrame(main_frame, text="초기 설정", padding="6")
        input_frame.pack(fill=tk.X, pady=(0, 6))

        # 결과 저장 폴더 설정
        ttk.Label(input_frame, text="결과 저장 폴더:").pack(anchor=tk.W)
        output_row = ttk.Frame(input_frame)
        output_row.pack(fill=tk.X, pady=(0, 5))

        self.output_dir_var = tk.StringVar(value=load_last_output_dir())
        output_entry = ttk.Entry(output_row, textvariable=self.output_dir_var)
        output_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

        self.browse_output_button = ttk.Button(output_row, text="찾아보기...",
                                              command=self.browse_output_dir)
        self.browse_output_button.pack(side=tk.LEFT, padx=(0, 5))

        ttk.Button(output_row, text="폴더 열기",
                  command=self.open_output_dir).pack(side=tk.LEFT)

        # 시드 키워드 입력 (한 줄에 하나씩, 실행 중에도 추가 가능)
        ttk.Label(input_frame, text="키워드 입력 (한 줄에 하나씩 입력, 실행 중에도 추가 가능):").pack(anchor=tk.W)
        seed_row = ttk.Frame(input_frame)
        seed_row.pack(fill=tk.X, pady=(0, 5))

        self.seed_keywords_text = tk.Text(seed_row, height=2, width=70)
        self.seed_keywords_text.pack(side=tk.LEFT, fill=tk.X, expand=True)

        seed_button_frame = ttk.Frame(seed_row)
        seed_button_frame.pack(side=tk.LEFT, padx=(10, 0), fill=tk.Y)
        ttk.Button(seed_button_frame, text="키워드 추가",
                  command=self.add_keywords_to_queue).pack(fill=tk.X)

        # 필수 키워드 입력
        ttk.Label(input_frame, text="필수 키워드 (쉼표로 구분, 선택사항):").pack(anchor=tk.W)
        self.required_keywords_var = tk.StringVar()
        required_entry = ttk.Entry(input_frame, textvariable=self.required_keywords_var, width=80)
        required_entry.pack(fill=tk.X)

        # 실행 버튼 영역 (입력 바로 아래, 항상 보이도록 배치)
        button_frame = ttk.LabelFrame(main_frame, text="실행", padding="6")
        button_frame.pack(fill=tk.X, pady=(0, 6))

        # 통합 실행 버튼들
        integrated_frame = ttk.Frame(button_frame)
        integrated_frame.pack(fill=tk.X, pady=(0, 5))

        self.start_button = ttk.Button(integrated_frame, text="전체 분석 시작",
                                     command=self.start_analysis, style="Accent.TButton")
        self.start_button.pack(side=tk.LEFT, padx=(0, 5))

        self.partial_button = ttk.Button(integrated_frame, text="A~F 정렬 분석",
                                       command=self.start_partial_analysis, style="Accent.TButton")
        self.partial_button.pack(side=tk.LEFT, padx=(0, 5))

        self.stop_button = ttk.Button(integrated_frame, text="중단",
                                    command=self.stop_analysis, state=tk.DISABLED)
        self.stop_button.pack(side=tk.LEFT)

        # 개별 실행 버튼들 (한 줄로 배치)
        individual_frame = ttk.Frame(button_frame)
        individual_frame.pack(fill=tk.X)

        ttk.Label(individual_frame, text="개별 모듈:").pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(individual_frame, text="A (우측연관)", width=11,
                  command=lambda: self.run_individual_module('A')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="B (자동완성)", width=11,
                  command=lambda: self.run_individual_module('B')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="C (키워드통합)", width=12,
                  command=lambda: self.run_individual_module('C')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="D (검색결과)", width=11,
                  command=lambda: self.run_individual_module('D')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="E (필터링)", width=10,
                  command=lambda: self.run_individual_module('E')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="F (최근데이터)", width=12,
                  command=lambda: self.run_individual_module('F')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="F_sort (정렬)", width=11,
                  command=lambda: self.run_individual_module('F_sort')).pack(side=tk.LEFT, padx=1)
        ttk.Button(individual_frame, text="G (블로거)", width=10,
                  command=lambda: self.run_individual_module('G')).pack(side=tk.LEFT, padx=1)

        # 버튼 스타일 설정
        style = ttk.Style()
        style.configure("Accent.TButton", font=("Arial", 10, "bold"))

        # 진행률 표시 (좌우로 나눠 배치하여 세로 공간을 절약)
        progress_frame = ttk.LabelFrame(main_frame, text="진행 상황", padding="6")
        progress_frame.pack(fill=tk.X, pady=(0, 6))

        overall_col = ttk.Frame(progress_frame)
        overall_col.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 10))

        current_col = ttk.Frame(progress_frame)
        current_col.pack(side=tk.LEFT, fill=tk.X, expand=True)

        # 전체 키워드 진행률 (대기열 전체 기준: 몇 개 중 몇 개 완료했는지)
        ttk.Label(overall_col, text="전체 키워드 진행").pack(anchor=tk.W)
        self.overall_progress_var = tk.DoubleVar()
        self.overall_progress_bar = ttk.Progressbar(overall_col, variable=self.overall_progress_var,
                                                   maximum=100, mode='determinate')
        self.overall_progress_bar.pack(fill=tk.X, pady=(0, 2))

        self.overall_status_label = ttk.Label(overall_col, text="대기 중인 키워드가 없습니다.")
        self.overall_status_label.pack(anchor=tk.W)

        # 현재 키워드의 단계(A~G) 진행률
        ttk.Label(current_col, text="현재 키워드 단계 진행").pack(anchor=tk.W)
        self.progress_var = tk.DoubleVar()
        self.progress_bar = ttk.Progressbar(current_col, variable=self.progress_var,
                                          maximum=100, mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=(0, 2))

        # 진행 상태 레이블 + 현재 단계 표시를 한 줄로 배치
        status_row = ttk.Frame(current_col)
        status_row.pack(fill=tk.X)
        self.status_label = ttk.Label(status_row, text="준비 완료")
        self.status_label.pack(side=tk.LEFT)
        self.current_step_label = ttk.Label(status_row, text="")
        self.current_step_label.pack(side=tk.LEFT, padx=(10, 0))

        # 아래쪽 영역(키워드 현황 + 로그)을 좌우로 나눠 갖도록 PanedWindow 사용 (로그를 오른쪽에 크게 표시)
        lower_pane = ttk.PanedWindow(main_frame, orient=tk.HORIZONTAL)
        lower_pane.pack(fill=tk.BOTH, expand=True)

        # 키워드 처리 현황 (왼쪽)
        keyword_status_frame = ttk.LabelFrame(lower_pane, text="키워드 처리 현황", padding="6")
        lower_pane.add(keyword_status_frame, weight=1)

        # 키워드 삭제 / 우선순위 변경 버튼 (먼저 pack하여 하단에 항상 고정 표시 - 리스트가 확장되어도 가려지지 않음)
        tree_button_frame = ttk.Frame(keyword_status_frame)
        tree_button_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=(5, 0))

        ttk.Button(tree_button_frame, text="선택 삭제",
                  command=self.delete_selected_keywords).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(tree_button_frame, text="▲ 위로",
                  command=lambda: self.move_selected_keyword(-1)).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(tree_button_frame, text="▼ 아래로",
                  command=lambda: self.move_selected_keyword(1)).pack(side=tk.LEFT)

        tree_container = ttk.Frame(keyword_status_frame)
        tree_container.pack(fill=tk.BOTH, expand=True)

        self.keyword_tree = ttk.Treeview(tree_container, columns=("keyword", "status"),
                                        show="headings", height=5, selectmode="extended")
        self.keyword_tree.heading("keyword", text="키워드")
        self.keyword_tree.heading("status", text="상태")
        self.keyword_tree.column("keyword", width=200)
        self.keyword_tree.column("status", width=90, anchor=tk.CENTER)
        self.keyword_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        tree_scroll = ttk.Scrollbar(tree_container, orient=tk.VERTICAL,
                                   command=self.keyword_tree.yview)
        tree_scroll.pack(side=tk.LEFT, fill=tk.Y)
        self.keyword_tree.configure(yscrollcommand=tree_scroll.set)

        # 우클릭 컨텍스트 메뉴 (삭제 / 순서 변경)
        self.keyword_context_menu = tk.Menu(self.keyword_tree, tearoff=0)
        self.keyword_context_menu.add_command(label="삭제", command=self.delete_selected_keywords)
        self.keyword_context_menu.add_command(label="위로 이동", command=lambda: self.move_selected_keyword(-1))
        self.keyword_context_menu.add_command(label="아래로 이동", command=lambda: self.move_selected_keyword(1))
        self.keyword_tree.bind("<Button-3>", self.show_keyword_context_menu)
        self.keyword_tree.bind("<Delete>", lambda e: self.delete_selected_keywords())

        # 로그 표시 영역 (오른쪽, 넓게)
        log_frame = ttk.LabelFrame(lower_pane, text="실행 로그", padding="6")
        lower_pane.add(log_frame, weight=2)

        # 스크롤 가능한 텍스트 위젯
        self.log_text = scrolledtext.ScrolledText(log_frame, wrap=tk.WORD, height=10,
                                                 font=("Consolas", 9))
        self.log_text.pack(fill=tk.BOTH, expand=True)

    def browse_output_dir(self):
        """결과 저장 폴더를 선택하고, 다음 실행을 위해 저장"""
        initial = self.output_dir_var.get().strip() or os.getcwd()
        if not os.path.isdir(initial):
            initial = os.getcwd()
        selected = filedialog.askdirectory(initialdir=initial, title="결과 저장 폴더 선택")
        if selected:
            self.output_dir_var.set(selected)
            save_last_output_dir(selected)

    def open_output_dir(self):
        """결과 저장 폴더를 탐색기로 연다"""
        path = self.output_dir_var.get().strip()
        if not path:
            messagebox.showwarning("경고", "결과 저장 폴더가 지정되지 않았습니다.")
            return

        try:
            os.makedirs(path, exist_ok=True)
            os.startfile(path)
        except Exception as e:
            messagebox.showerror("오류", f"폴더를 열 수 없습니다:\n{e}")

    def apply_output_dir(self) -> bool:
        """결과 저장 폴더를 생성/이동(chdir)하고 마지막 사용 폴더로 저장. 성공 시 True"""
        path = self.output_dir_var.get().strip()
        if not path:
            messagebox.showerror("입력 오류", "결과 저장 폴더를 지정해주세요.")
            return False

        try:
            os.makedirs(path, exist_ok=True)
            os.chdir(path)
            save_last_output_dir(path)
            return True
        except Exception as e:
            messagebox.showerror("오류", f"결과 저장 폴더로 이동할 수 없습니다:\n{e}")
            return False

    def parse_keywords(self, keyword_string: str) -> List[str]:
        """키워드 문자열을 리스트로 변환 (줄바꿈 또는 쉼표로 구분)"""
        if not keyword_string or not keyword_string.strip():
            return []

        keywords = []
        for line in keyword_string.splitlines():
            for kw in line.split(","):
                kw = kw.strip()
                if kw:
                    keywords.append(kw)
        return keywords

    def get_seed_keywords_input(self) -> List[str]:
        """키워드 입력창의 내용을 리스트로 변환"""
        return self.parse_keywords(self.seed_keywords_text.get("1.0", tk.END))

    def enqueue_keyword(self, keyword: str):
        """키워드를 처리 대기열과 현황 트리에 추가"""
        self._keyword_counter += 1
        iid = f"kw_{self._keyword_counter}"
        self.keyword_tree.insert("", tk.END, iid=iid, values=(keyword, "대기중"))
        self.resync_queue_from_tree()
        self.save_tasks()
        self.update_overall_progress()

    def add_keywords_to_queue(self):
        """입력창의 키워드들을 대기열에 추가 (실행 중에도 사용 가능)"""
        keywords = self.get_seed_keywords_input()
        if not keywords:
            messagebox.showwarning("입력 오류", "추가할 키워드를 입력해주세요.")
            return

        for kw in keywords:
            self.enqueue_keyword(kw)

        self.seed_keywords_text.delete("1.0", tk.END)

        if self.is_running:
            print(f"{len(keywords)}개의 키워드가 처리 대기열에 추가되었습니다: {keywords}")

    def set_keyword_status(self, iid: str, status: str):
        """키워드 현황 트리의 상태를 스레드 안전하게 갱신"""
        self.root.after(0, lambda: self._set_keyword_status_ui(iid, status))

    def _set_keyword_status_ui(self, iid: str, status: str):
        if self.keyword_tree.exists(iid):
            self.keyword_tree.set(iid, "status", status)
            self.save_tasks()
            self.update_overall_progress()

    def update_overall_progress(self):
        """대기열 전체 키워드 기준 진행률(완료/처리중/대기/실패 개수)을 진행 상황 영역에 갱신.

        키워드 하나당 A~G 세부 단계 진행률만으로는 "전체 중 몇 개나 남았는지"를 알 수 없어
        추가한 상단 진행 표시. 트리(키워드 처리 현황)를 소스로 삼아 매번 다시 집계한다.
        """
        total = 0
        done = 0
        failed = 0
        processing = 0
        waiting = 0

        for iid in self.keyword_tree.get_children():
            status = self.keyword_tree.set(iid, "status")
            total += 1
            if status == "완료 ✓":
                done += 1
            elif status == "실패 ✗":
                failed += 1
            elif status == "처리중":
                processing += 1
            elif status == "대기중":
                waiting += 1
            # "중단됨" 등 그 외 상태는 done/failed/processing/waiting 어디에도 세지 않되 total에는 포함

        if total == 0:
            self.overall_progress_var.set(0)
            self.overall_status_label.config(text="대기 중인 키워드가 없습니다.")
            return

        finished = done + failed
        self.overall_progress_var.set((finished / total) * 100)
        self.overall_status_label.config(
            text=(f"전체 {total}개 중 {finished}개 처리 완료 "
                  f"(성공 {done} / 실패 {failed}) · 처리중 {processing} · 남은 키워드 {waiting}개")
        )

    def resync_queue_from_tree(self):
        """트리에 표시된 순서/상태를 기준으로 대기열 내부 순서를 다시 구성 ('대기중' 항목만 대상).

        처리 중인 항목(self.processing_iid)은 상태 갱신이 반영되기 전 짧은 순간에도
        트리에는 여전히 '대기중'으로 보일 수 있으므로 항상 제외한다 (중복 처리 방지).
        """
        pending = []
        for iid in self.keyword_tree.get_children():
            if iid == self.processing_iid:
                continue
            if self.keyword_tree.set(iid, "status") == "대기중":
                keyword = self.keyword_tree.set(iid, "keyword")
                pending.append((iid, keyword))
        self.keyword_queue.replace_all(pending)

    def save_tasks(self):
        """현재 키워드 처리 현황을 파일로 저장 (재시작 시 이어서 진행하기 위함)"""
        try:
            tasks = []
            for iid in self.keyword_tree.get_children():
                keyword = self.keyword_tree.set(iid, "keyword")
                status = self.keyword_tree.set(iid, "status")
                tasks.append({"iid": iid, "keyword": keyword, "status": status})
            with open(TASKS_FILE, "w", encoding="utf-8") as f:
                json.dump(tasks, f, ensure_ascii=False, indent=2)
        except Exception as e:
            print(f"작업 목록 저장 실패: {e}")

    def load_tasks(self):
        """저장된 키워드 처리 현황을 불러와 복원. 처리 중이던/중단된 항목은 대기중으로 되돌려 이어서 처리"""
        try:
            with open(TASKS_FILE, "r", encoding="utf-8") as f:
                tasks = json.load(f)
        except FileNotFoundError:
            return
        except Exception as e:
            print(f"작업 목록 불러오기 실패: {e}")
            return

        if not tasks:
            return

        max_counter = 0
        resumed_count = 0

        for task in tasks:
            iid = task.get("iid", "")
            keyword = task.get("keyword", "")
            status = task.get("status", "대기중")
            if not iid or not keyword or self.keyword_tree.exists(iid):
                continue

            # 프로그램이 종료될 당시 처리 중이었거나 중단된 항목은 완료되지 않았으므로 다시 대기시킨다
            if status in ("처리중", "중단됨"):
                status = "대기중"
            if status == "대기중":
                resumed_count += 1

            self.keyword_tree.insert("", tk.END, iid=iid, values=(keyword, status))

            try:
                n = int(iid.split("_", 1)[1])
                max_counter = max(max_counter, n)
            except (IndexError, ValueError):
                pass

        self._keyword_counter = max_counter
        self.resync_queue_from_tree()
        self.update_overall_progress()

        if resumed_count:
            print(f"이전 실행에서 이어서 처리할 키워드 {resumed_count}개를 불러왔습니다. '전체 분석 시작'을 누르면 이어서 진행됩니다.")

    def delete_selected_keywords(self):
        """선택한 키워드를 현황 목록과 대기열에서 삭제 (처리 중인 항목은 삭제 불가)"""
        selected = list(self.keyword_tree.selection())
        if not selected:
            messagebox.showinfo("안내", "삭제할 키워드를 선택해주세요.")
            return

        blocked = [iid for iid in selected
                   if iid == self.processing_iid or self.keyword_tree.set(iid, "status") == "처리중"]
        deletable = [iid for iid in selected if iid not in blocked]

        for iid in deletable:
            if self.keyword_tree.exists(iid):
                self.keyword_tree.delete(iid)

        if deletable:
            self.resync_queue_from_tree()
            self.save_tasks()
            self.update_overall_progress()

        if blocked:
            messagebox.showwarning("삭제 불가", "처리 중인 키워드는 삭제할 수 없습니다. 완료를 기다리거나 먼저 중단해주세요.")

    def move_selected_keyword(self, direction: int):
        """선택한 키워드(들)의 우선순위를 함께 변경 (direction: -1=위로, 1=아래로).
        여러 개 선택 시 상대적인 순서를 유지한 채 블록으로 이동한다. 대기중 상태만 가능"""
        selected = set(self.keyword_tree.selection())
        if not selected:
            messagebox.showinfo("안내", "순서를 변경할 키워드를 선택해주세요.")
            return

        if self.processing_iid in selected:
            messagebox.showwarning("이동 불가", "처리 중인 키워드는 순서를 변경할 수 없습니다.")
            return

        if any(self.keyword_tree.set(iid, "status") != "대기중" for iid in selected):
            messagebox.showwarning("이동 불가", "대기중인 키워드만 순서를 변경할 수 있습니다.")
            return

        children = list(self.keyword_tree.get_children())
        indices = sorted(children.index(iid) for iid in selected)
        # 위로 이동은 위쪽(작은 인덱스)부터, 아래로 이동은 아래쪽(큰 인덱스)부터 처리해야
        # 블록 내부 순서가 무너지지 않는다.
        order = indices if direction < 0 else list(reversed(indices))

        moved = False
        for idx in order:
            iid = children[idx]
            new_idx = idx + direction
            if new_idx < 0 or new_idx >= len(children):
                continue

            neighbor_iid = children[new_idx]
            if neighbor_iid in selected:
                # 이미 선택된(같은 블록) 항목이면 함께 이동 중이므로 건너뜀
                continue
            if self.keyword_tree.set(neighbor_iid, "status") != "대기중":
                # 대기중이 아닌 항목(처리중/완료/실패 등)은 건너뛰지 않음
                continue

            self.keyword_tree.move(iid, "", new_idx)
            children[idx], children[new_idx] = children[new_idx], children[idx]
            moved = True

        if moved:
            self.resync_queue_from_tree()
            self.save_tasks()

    def show_keyword_context_menu(self, event):
        """키워드 트리 우클릭 시 삭제/순서 변경 메뉴 표시"""
        iid = self.keyword_tree.identify_row(event.y)
        if iid:
            if iid not in self.keyword_tree.selection():
                self.keyword_tree.selection_set(iid)
            self.keyword_context_menu.tk_popup(event.x_root, event.y_root)

    def start_partial_analysis(self):
        """A~F 정렬 단계만 분석 시작 (지속 실행되는 대기열 소비자 시작)"""
        self._start_queue_consumer(include_g=False)

    def start_analysis(self):
        """분석 시작 (지속 실행되는 대기열 소비자 시작)"""
        self._start_queue_consumer(include_g=True)

    def _start_queue_consumer(self, include_g: bool):
        # 입력창에 남아있는 키워드를 대기열에 편입
        pending_keywords = self.get_seed_keywords_input()
        for kw in pending_keywords:
            self.enqueue_keyword(kw)
        if pending_keywords:
            self.seed_keywords_text.delete("1.0", tk.END)

        if self.is_running:
            # 이미 실행 중이면 새로 입력한 키워드는 대기열에 자동으로 추가되어 처리됨
            if pending_keywords:
                print(f"{len(pending_keywords)}개의 키워드가 처리 대기열에 추가되었습니다: {pending_keywords}")
            return

        if self.keyword_queue.empty():
            messagebox.showerror("입력 오류", "처리할 키워드를 최소 1개 이상 입력해주세요.")
            return

        # 결과 저장 폴더로 이동 (없으면 생성) 및 마지막 사용 폴더로 저장
        if not self.apply_output_dir():
            return

        # 기존 결과 파일 삭제
        result_file = "result/keywordList_all.xlsx"
        if os.path.exists(result_file):
            try:
                os.remove(result_file)
                print(f"기존 결과 파일 삭제됨: {result_file}")
            except Exception as e:
                print(f"기존 결과 파일 삭제 실패: {e}")

        # 상태 변경
        self.is_running = True
        self.stop_requested = False
        self.start_button.config(state=tk.DISABLED)
        self.partial_button.config(state=tk.DISABLED)
        self.stop_button.config(state=tk.NORMAL)
        self.browse_output_button.config(state=tk.DISABLED)

        required_keywords = self.parse_keywords(self.required_keywords_var.get())

        # 백그라운드에서 대기열 소비자 실행 (실행 중 계속 추가되는 키워드도 처리)
        consumer_thread = threading.Thread(target=self.run_queue_consumer,
                                          args=(required_keywords, include_g))
        consumer_thread.daemon = True
        consumer_thread.start()

    def stop_analysis(self):
        """분석 중단"""
        if not self.is_running:
            return

        self.stop_requested = True
        self.status_label.config(text="중단 요청됨")

    def run_individual_module(self, module_name: str):
        """개별 모듈 실행"""
        if self.is_running:
            messagebox.showwarning("경고", "현재 다른 작업이 실행 중입니다.")
            return

        # 입력 검증
        seed_keywords = self.get_seed_keywords_input()
        required_keywords = self.parse_keywords(self.required_keywords_var.get())

        # 모듈별 입력 요구사항 확인
        if module_name in ['A', 'B', 'G'] and not seed_keywords:
            messagebox.showerror("입력 오류", f"{module_name} 모듈은 시드 키워드가 필요합니다.")
            return
        if module_name == 'C' and not required_keywords:
            messagebox.showwarning("입력 확인", "C 모듈에 필수 키워드가 설정되지 않았습니다.\n모든 데이터를 유지합니다.")

        # 결과 저장 폴더로 이동 (없으면 생성) 및 마지막 사용 폴더로 저장
        if not self.apply_output_dir():
            return

        # 실행 시작
        self.is_running = True
        self.stop_requested = False
        self.start_button.config(state=tk.DISABLED)
        self.stop_button.config(state=tk.NORMAL)
        self.browse_output_button.config(state=tk.DISABLED)

        # 백그라운드에서 모듈 실행
        module_thread = threading.Thread(target=self.run_single_module,
                                       args=(module_name, seed_keywords, required_keywords))
        module_thread.daemon = True
        module_thread.start()

    def run_single_module(self, module_name: str, seed_keywords: List[str], required_keywords: List[str]):
        """단일 모듈 실행"""
        try:
            module_map = {
                'A': ('A_rightside', lambda: self.run_module_a(seed_keywords[0] if seed_keywords else "")),
                'B': ('B_autocomplete', lambda: self.run_module_b(seed_keywords[0] if seed_keywords else "")),
                'C': ('C_sumKeyword', lambda: self.run_module_c(required_keywords)),
                'D': ('D_searchresult', lambda: self.run_module_d()),
                'E': ('E_deleteAndPriority', lambda: self.run_module_e()),
                'F': ('F_add_recent30days', lambda: self.run_module_f()),
                'F_sort': ('F_sort_result_1', lambda: self.run_module_f_sort()),
                'G': ('G_add_blogger_by_mainPage', lambda: self.run_module_g(seed_keywords[0] if seed_keywords else ""))
            }

            module_title, module_func = module_map[module_name]

            self.update_progress(0, 1, f"{module_name}단계: {module_title} 실행중")

            success = module_func()

            if success:
                self.update_progress(1, 1, f"{module_name}단계 완료")

                success_msg = f"{module_name} 모듈이 성공적으로 실행되었습니다!"
                messagebox.showinfo("완료", success_msg)
            else:
                messagebox.showerror("실패", f"{module_name} 모듈 실행에 실패했습니다.")

        except Exception as e:
            messagebox.showerror("오류", f"{module_name} 모듈 실행 중 오류가 발생했습니다:\n{str(e)}")

        finally:
            # UI 상태 초기화
            self.root.after(0, self.reset_ui)

    def process_single_keyword(self, seed_keyword: str, required_keywords: List[str], include_g: bool) -> bool:
        """단일 키워드에 대해 A~F_sort(또는 A~G) 단계를 순차 실행"""
        steps = [
            ("A단계: 우측 연관 검색어 추출", lambda: self.run_module_a(seed_keyword)),
            ("B단계: 자동완성 검색어 추출", lambda: self.run_module_b(seed_keyword)),
            ("C단계: 키워드 데이터 통합", lambda: self.run_module_c(required_keywords)),
            ("D단계: 검색 결과 처리", lambda: self.run_module_d()),
            ("E단계: 데이터 필터링 및 우선순위", lambda: self.run_module_e()),
            ("F단계: 최근 30일 블로그 데이터 추가", lambda: self.run_module_f()),
            ("F_sort단계: 최근 30일 데이터 정렬 및 분석", lambda: self.run_module_f_sort()),
        ]
        if include_g:
            steps.append(("G단계: 블로거 정보 관리", lambda: self.run_module_g(seed_keyword)))

        total_steps = len(steps)
        for step_idx, (step_label, step_func) in enumerate(steps, start=1):
            if self.stop_requested:
                return False
            self.update_progress(step_idx, total_steps, f"'{seed_keyword}' - {step_label}")
            if not step_func():
                return False

        # G단계까지 끝나면 결과 파일이 keywordResult_{seed}.xlsx로 이름이 바뀌므로 그에 맞춰 경로를 잡는다
        excel_path = f"keywordResult_{seed_keyword}.xlsx" if include_g else "result/keywordList_all.xlsx"
        self.report_golden_keywords(seed_keyword, excel_path)

        return True

    def report_golden_keywords(self, seed_keyword: str, excel_path: str):
        """분석 완료 후, 최근 30일 블로그 발행수가 GOLDEN_BLOG_THRESHOLD 미만인 키워드를
        최대 GOLDEN_TARGET_COUNT개까지 찾아 실행 로그 창에 한 줄에 하나씩 출력한다.
        (엑셀을 열어보지 않아도 결과를 바로 확인할 수 있도록 하기 위함)

        1차로 F_sort 단계에서 이미 계산된 recent30days_sorted(없으면 recent30days) 시트에서 찾고,
        그래도 부족하면 검색광고 API에서 검색량 상위 51~100위 후보를 추가로 블로그수 확인해 채운다.
        """
        try:
            if not os.path.exists(excel_path):
                print(f"'{seed_keyword}' - 결과 파일을 찾을 수 없어 황금키워드 선별을 건너뜁니다: {excel_path}")
                return

            wb = load_workbook(excel_path, data_only=True)
            sheet_name = "recent30days_sorted" if "recent30days_sorted" in wb.sheetnames else "recent30days"
            if sheet_name not in wb.sheetnames:
                print(f"'{seed_keyword}' - {sheet_name} 시트가 없어 황금키워드 선별을 건너뜁니다.")
                return

            ws = wb[sheet_name]
            # recent30days(_sorted) 컬럼: A=seed_keyword, B=relKeyword, C=monthlyTotal, D=recent30dayblog
            candidates = {}  # keyword -> (search_qc, blog_count)
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or len(row) < 4:
                    continue
                kw, qc, blog = row[1], row[2], row[3]
                if not kw:
                    continue
                try:
                    blog_n = int(blog)
                    qc_n = int(qc) if qc is not None else 0
                except (TypeError, ValueError):
                    continue
                if blog_n < GOLDEN_BLOG_THRESHOLD:
                    prev = candidates.get(kw)
                    if prev is None or qc_n > prev[0]:
                        candidates[kw] = (qc_n, blog_n)

            golden = sorted(candidates.items(), key=lambda x: x[1][0], reverse=True)

            # 부족하면 검색광고 API 상위 51~100위 후보를 추가로 블로그수 확인해서 채운다
            if len(golden) < GOLDEN_TARGET_COUNT:
                client_id = os.getenv("NAVER_CLIENT_ID", "").strip()
                client_secret = os.getenv("NAVER_CLIENT_SECRET", "").strip()
                if client_id and client_secret:
                    print(f"'{seed_keyword}' - 최종 결과에서 {len(golden)}개 확보. "
                          f"검색광고 API 상위 51~100위 후보를 추가로 확인합니다...")
                    try:
                        data = A_rightside.call_keywordstool(seed_keyword, show_detail=1)
                        keyword_list = data.get("keywordList") or []
                        seen_kw = set(candidates.keys()) | {seed_keyword}
                        extra_rows = []
                        for r in keyword_list:
                            kw = str(r.get("relKeyword", "")).strip()
                            if not kw or kw in seen_kw:
                                continue
                            try:
                                qc_n = int(r.get("monthlyPcQcCnt") or 0) + int(r.get("monthlyMobileQcCnt") or 0)
                            except (TypeError, ValueError):
                                qc_n = 0
                            extra_rows.append((kw, qc_n))

                        extra_rows.sort(key=lambda x: x[1], reverse=True)
                        extra_pool = extra_rows[50:100]  # 검색량 상위 51~100위 후보

                        for kw, qc_n in extra_pool:
                            if len(golden) >= GOLDEN_TARGET_COUNT:
                                break
                            blog_n = A_rightside._get_recent_30day_blog_count(kw, client_id, client_secret)
                            if blog_n is not None and blog_n < GOLDEN_BLOG_THRESHOLD:
                                golden.append((kw, (qc_n, blog_n)))
                            time.sleep(0.1)

                        golden.sort(key=lambda x: x[1][0], reverse=True)
                    except Exception as e:
                        print(f"'{seed_keyword}' - 추가 후보 조회 중 오류 발생: {e}")

            golden = golden[:GOLDEN_TARGET_COUNT]

            print(f"\n=== '{seed_keyword}' 황금키워드 선별 결과 "
                  f"(최근 30일 블로그 {GOLDEN_BLOG_THRESHOLD}개 미만, {len(golden)}개) ===")
            for kw, (qc_n, blog_n) in golden:
                print(f"{kw}\t검색량 {qc_n}\t블로그(30일) {blog_n}")
            if len(golden) < GOLDEN_TARGET_COUNT:
                print(f"(후보가 부족해 {len(golden)}개만 찾았습니다)")

        except Exception as e:
            print(f"'{seed_keyword}' - 황금키워드 선별 중 오류 발생: {e}")

    def run_queue_consumer(self, required_keywords: List[str], include_g: bool):
        """대기열에서 키워드를 꺼내 순차 처리. 실행 중 새로 추가되는 키워드도 계속 소비한다."""
        try:
            while not self.stop_requested:
                try:
                    iid, seed_keyword = self.keyword_queue.get(timeout=1)
                except queue.Empty:
                    continue

                # resync_queue_from_tree()가 처리 중인 항목을 다시 대기열에 넣지 않도록 표시
                self.processing_iid = iid
                self.set_keyword_status(iid, "처리중")

                try:
                    success = self.process_single_keyword(seed_keyword, required_keywords, include_g)
                except Exception as e:
                    success = False
                    print(f"'{seed_keyword}' 처리 중 오류가 발생했습니다: {e}")

                if success:
                    self.set_keyword_status(iid, "완료 ✓")
                    result_note = "keywordResult_{}.xlsx".format(seed_keyword) if include_g else "F_sort 단계까지"
                    print(f"'{seed_keyword}' 키워드 처리가 완료되었습니다. ({result_note})")
                elif self.stop_requested:
                    self.set_keyword_status(iid, "중단됨")
                else:
                    self.set_keyword_status(iid, "실패 ✗")
                    print(f"'{seed_keyword}' 키워드 처리에 실패했습니다.")

                self.processing_iid = None
                self.keyword_queue.task_done()

        finally:
            self.processing_iid = None
            self.root.after(0, self.reset_ui)

    def update_progress(self, current: int, total: int, status: str):
        """진행률 업데이트"""
        progress = (current / total) * 100
        self.root.after(0, lambda: self.progress_var.set(progress))
        self.root.after(0, lambda: self.status_label.config(text=f"{current}/{total} 단계 진행중"))
        self.root.after(0, lambda: self.current_step_label.config(text=status))

    def run_module_a(self, seed_keyword: str) -> bool:
        """A 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            # monkey patch로 입력 우회
            original_input = __builtins__.input
            mock_input = MockInput(seed_keyword)
            __builtins__.input = mock_input

            try:
                A_rightside.main()
                return True
            finally:
                __builtins__.input = original_input

        except KeyboardInterrupt:
            return False
        except Exception as e:
            return False

    def run_module_b(self, seed_keyword: str, max_retries: int = 3) -> bool:
        """B 모듈 실행

        자동완성 차단(AutoCompleteBlockedError)이나 그 외 오류가 발생해도 팝업을 띄우지 않고,
        스스로 대기한 뒤 브라우저를 껐다 켜서(새 세션) 같은 키워드로 최대 max_retries회 재시도한다.
        모두 실패하면 이 키워드만 실패 처리하고 다음 키워드로 넘어간다(전체 실행은 중단하지 않음).
        """
        if self.stop_requested:
            return False

        # Selenium 사용 가능 여부 확인
        if not hasattr(B_autocomplete, 'SELENIUM_AVAILABLE') or not B_autocomplete.SELENIUM_AVAILABLE:
            return True  # 실패로 처리하지 않고 건너뛰기

        for attempt in range(1, max_retries + 1):
            if self.stop_requested:
                return False

            original_input = __builtins__.input
            mock_input = MockInput(seed_keyword)
            __builtins__.input = mock_input

            try:
                B_autocomplete.main()
                return True
            except KeyboardInterrupt:
                return False
            except B_autocomplete.AutoCompleteBlockedError:
                print(f"'{seed_keyword}' - 자동완성 차단이 감지되었습니다 ({attempt}/{max_retries}회차).")
            except Exception as e:
                print(f"'{seed_keyword}' - B단계(자동완성) 처리 중 오류 발생 ({attempt}/{max_retries}회차): {e}")
            finally:
                __builtins__.input = original_input

            # 브라우저를 새로 띄워 재시도하므로, 이전 세션에서 누적된 연속 실패 카운트를 초기화
            B_autocomplete.AutoCompleteSession._consecutive_failures = 0

            if attempt < max_retries:
                wait_sec = random.uniform(10, 20)
                print(f"'{seed_keyword}' - {wait_sec:.0f}초 대기 후 브라우저를 재시작하여 재시도합니다. ({attempt}/{max_retries})")
                time.sleep(wait_sec)

        print(f"'{seed_keyword}' - B단계(자동완성) {max_retries}회 재시도 모두 실패하여 다음 키워드로 넘어갑니다.")
        return False

    def run_module_c(self, required_keywords: List[str]) -> bool:
        """C 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            # monkey patch로 입력 우회
            original_input = __builtins__.input
            mock_input = MockInput(required_keywords=required_keywords)
            __builtins__.input = mock_input

            try:
                C_sumKeyword.main()
                return True
            finally:
                __builtins__.input = original_input

        except KeyboardInterrupt:
            return False
        except Exception as e:
            return False

    def run_module_d(self) -> bool:
        """D 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            D_searchresult.main()
            return True
        except Exception as e:
            return False

    def run_module_e(self) -> bool:
        """E 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            E_deleteAndPriority.main()
            return True
        except Exception as e:
            return False

    def run_module_f(self) -> bool:
        """F 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            F_add_recent30days.main()
            return True
        except Exception as e:
            return False

    def run_module_f_sort(self) -> bool:
        """F_sort 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            F_sort_result_1.main()
            return True
        except Exception as e:
            return False

    def run_module_g(self, seed_keyword: str) -> bool:
        """G 모듈 실행"""
        if self.stop_requested:
            return False

        try:
            # 명령줄 인자로 시드 키워드 전달
            original_argv = sys.argv.copy()
            sys.argv = ["G_add_blogger_by_mainPage.py", seed_keyword]

            try:
                G_add_blogger_by_mainPage.main()

                # G 단계 완료 후 파일명 변경 및 기존 파일 삭제
                self.rename_and_cleanup_file(seed_keyword)

                return True
            finally:
                sys.argv = original_argv

        except Exception as e:
            return False

    def rename_and_cleanup_file(self, seed_keyword: str):
        """G 단계 완료 후 파일명 변경 및 정리.

        rename 실패(대상 파일이 Excel에서 열려있는 등)를 조용히 삼키면 GUI가 완료로
        표시되는데도 실제로는 결과 파일이 갱신되지 않는 문제가 있었으므로, 실패 시
        예외를 그대로 올려 호출자(run_module_g)가 실패로 처리하도록 한다.
        """
        excel_path = "result/keywordList_all.xlsx"
        new_filename = f"keywordResult_{seed_keyword}.xlsx"

        # 파일 존재 확인
        if not os.path.exists(excel_path):
            raise FileNotFoundError(f"{excel_path} 파일을 찾을 수 없습니다.")

        try:
            # 새 파일명으로 변경
            os.rename(excel_path, new_filename)
            print(f"파일명이 '{new_filename}'으로 변경되었습니다.")
        except OSError as e:
            raise RuntimeError(
                f"'{new_filename}' 파일명 변경 실패 (파일이 Excel 등에서 열려있지 않은지 확인하세요): {e}"
            ) from e


    def reset_ui(self):
        """UI 상태 초기화"""
        self.is_running = False
        self.stop_requested = False
        self.start_button.config(state=tk.NORMAL)
        self.partial_button.config(state=tk.NORMAL)
        self.stop_button.config(state=tk.DISABLED)
        self.browse_output_button.config(state=tk.NORMAL)
        self.status_label.config(text="준비 완료")
        self.current_step_label.config(text="")


def main():
    """메인 함수"""
    root = tk.Tk()
    app = IntegrationGUI(root)

    # 창이 닫힐 때 sys.stdout 복원
    def on_closing():
        if hasattr(app, 'original_stdout'):
            sys.stdout = app.original_stdout
        root.destroy()

    root.protocol("WM_DELETE_WINDOW", on_closing)
    root.mainloop()

    # mainloop 종료 후 sys.stdout 복원 (안전장치)
    if hasattr(app, 'original_stdout'):
        sys.stdout = app.original_stdout


if __name__ == "__main__":
    main()