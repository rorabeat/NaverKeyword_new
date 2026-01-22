from openpyxl import load_workbook

wb = load_workbook('result/keywordList_all.xlsx')
ws = wb['bloggerlist']

print('bloggerlist 시트 내용:')
print(f'총 행 수: {ws.max_row}')
print(f'총 열 수: {ws.max_column}')

# 헤더 출력
header = [ws.cell(row=1, column=col).value for col in range(1, ws.max_column + 1)]
print('헤더:', header)

# 데이터 행 출력 (최대 3행)
for row in range(2, min(5, ws.max_row + 1)):
    row_data = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
    print(f'행 {row}:', row_data)