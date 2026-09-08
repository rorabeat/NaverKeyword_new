import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox
import threading
import sys
import os
from datetime import datetime, timezone, timedelta

# 부모 디렉토리를 sys.path에 추가하여 import 가능하게 함
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# blogger_by_mainPage.py의 함수들 import
from blogger_by_mainPage import get_blogger_data_from_main_page
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment

class BloggerRankGUI:
    def __init__(self, root):
        self.root = root
        self.root.title("네이버 블로거 순위 조회 v1.0 - 메인페이지 검색 기반")
        self.root.geometry("800x600")
        self.root.resizable(True, True)

        # 환경 변수 로드
        self.load_env_variables()

        # 메인 프레임
        main_frame = ttk.Frame(root, padding="10")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        root.columnconfigure(0, weight=1)
        root.rowconfigure(0, weight=1)

        # 키워드 입력 섹션
        keyword_frame = ttk.LabelFrame(main_frame, text="키워드 입력", padding="5")
        keyword_frame.grid(row=0, column=0, sticky=(tk.W, tk.E), pady=(0, 10))
        keyword_frame.columnconfigure(1, weight=1)

        ttk.Label(keyword_frame, text="키워드:").grid(row=0, column=0, sticky=tk.W, padx=(0, 5))
        self.keyword_entry = ttk.Entry(keyword_frame, width=50)
        self.keyword_entry.grid(row=0, column=1, sticky=(tk.W, tk.E), padx=(0, 5))
        # 예시 텍스트 제거 - 사용자가 직접 입력하도록 유도

        ttk.Label(keyword_frame, text="(여러 개 입력시 컴마(,)로 구분)", font=("", 8)).grid(
            row=1, column=1, sticky=tk.W, padx=(0, 5), pady=(2, 0))

        # 실행 버튼과 폴더 열기 버튼
        button_frame = ttk.Frame(keyword_frame)
        button_frame.grid(row=0, column=2, rowspan=2, padx=(5, 0))

        self.run_button = ttk.Button(button_frame, text="실행", command=self.start_processing)
        self.run_button.pack(side=tk.TOP, pady=(0, 2))

        self.open_folder_button = ttk.Button(button_frame, text="결과 폴더 열기", command=self.open_result_folder)
        self.open_folder_button.pack(side=tk.TOP)

        # 로그 출력 섹션
        log_frame = ttk.LabelFrame(main_frame, text="진행 상황 로그", padding="5")
        log_frame.grid(row=1, column=0, sticky=(tk.W, tk.E, tk.N, tk.S), pady=(10, 0))
        log_frame.columnconfigure(0, weight=1)
        log_frame.rowconfigure(0, weight=1)

        # 스크롤 가능한 텍스트 박스
        self.log_text = scrolledtext.ScrolledText(log_frame, wrap=tk.WORD, height=20)
        self.log_text.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))

        # 메인 프레임 설정
        main_frame.columnconfigure(0, weight=1)
        main_frame.rowconfigure(1, weight=1)

        # 초기 로그 메시지
        self.log_message("=== 네이버 블로거 순위 조회 프로그램 ===")
        self.log_message("• 키워드를 입력하고 '실행' 버튼을 클릭하세요")
        self.log_message("• 여러 키워드는 컴마(,)로 구분하세요")
        self.log_message("• 각 키워드당 상위 5개 블로거 정보를 수집합니다")
        self.log_message("")
        self.check_env_status()

    def load_env_variables(self):
        """환경 변수 로드 및 검증"""
        try:
            from dotenv import load_dotenv
            load_dotenv()

            self.client_id = os.getenv("NAVER_CLIENT_ID", "").strip()
            self.client_secret = os.getenv("NAVER_CLIENT_SECRET", "").strip()

        except ImportError:
            self.log_message("경고: python-dotenv가 설치되지 않았습니다.")
            self.client_id = ""
            self.client_secret = ""

    def check_env_status(self):
        """환경 변수 상태 확인"""
        if not self.client_id or not self.client_secret:
            self.log_message("경고: NAVER_CLIENT_ID 또는 NAVER_CLIENT_SECRET이 설정되지 않았습니다.")
            self.log_message("프로그램이 정상 작동하지 않을 수 있습니다.")
        else:
            self.log_message("네이버 API 키가 정상적으로 로드되었습니다.")

    def open_result_folder(self):
        """결과 폴더 열기"""
        try:
            # 현재 소스코드 위치의 result 폴더 경로
            current_dir = os.path.dirname(os.path.abspath(__file__))
            result_dir = os.path.join(current_dir, "result")
            if os.path.exists(result_dir):
                os.startfile(result_dir)  # Windows에서 폴더 열기
                self.log_message("결과 폴더를 열었습니다.")
            else:
                self.log_message("결과 폴더가 존재하지 않습니다.")
        except Exception as e:
            self.log_message(f"결과 폴더 열기 실패: {str(e)}")

    def save_blogger_results_to_excel(self, keyword: str, run_timestamp: datetime, blogger_data: list) -> None:
        """블로거 검색 결과를 bloggerlevel.xlsx 파일에 누적 저장"""
        try:
            # 현재 소스코드 위치에 result 폴더 생성
            current_dir = os.path.dirname(os.path.abspath(__file__))
            result_dir = os.path.join(current_dir, "result")
            if not os.path.exists(result_dir):
                os.makedirs(result_dir)

            excel_path = f"{result_dir}/bloggerlevel.xlsx"

            # 헤더 정의 (항상 사용됨)
            headers = [
                "실행시각", "키워드",
                "1위 블로거", "1위 오늘 조회수", "1위 예상 일일 조회수",
                "2위 블로거", "2위 오늘 조회수", "2위 예상 일일 조회수",
                "3위 블로거", "3위 오늘 조회수", "3위 예상 일일 조회수",
                "4위 블로거", "4위 오늘 조회수", "4위 예상 일일 조회수",
                "5위 블로거", "5위 오늘 조회수", "5위 예상 일일 조회수"
            ]

            # 기존 파일이 있으면 로드, 없으면 새로 생성
            try:
                wb = load_workbook(excel_path)
                ws = wb.active
            except FileNotFoundError:
                wb = Workbook()
                ws = wb.active
                ws.title = "BloggerLevel"

                # 헤더 추가
                for col_idx, header in enumerate(headers, start=1):
                    cell = ws.cell(row=1, column=col_idx, value=header)
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # 필터 및 고정 설정
                ws.freeze_panes = "A2"
                ws.auto_filter.ref = f"A1:{chr(ord('A') + len(headers) - 1)}1"

            # 새로운 데이터를 첫 번째 행(헤더 다음)에 삽입하여 최신 데이터를 위쪽에 표시
            ws.insert_rows(2)

            # 데이터 준비
            row_data = [
                run_timestamp.strftime('%Y-%m-%d %H:%M:%S'),
                keyword
            ]

            # 블로거 데이터 추가 (최대 5개)
            for i in range(5):
                if i < len(blogger_data):
                    blogger_info = blogger_data[i]
                    row_data.extend([
                        blogger_info.get("blogger", ""),
                        blogger_info.get("today", ""),
                        blogger_info.get("est", "")
                    ])
                else:
                    row_data.extend(["", "", ""])

            # 데이터 입력
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=2, column=col_idx, value=value)

            # 컬럼 너비 자동 조정
            for col_idx in range(1, len(headers) + 1):
                max_len = 10
                for row_idx in range(1, ws.max_row + 1):
                    cell_value = ws.cell(row=row_idx, column=col_idx).value
                    if cell_value:
                        max_len = max(max_len, len(str(cell_value)))
                ws.column_dimensions[chr(ord('A') + col_idx - 1)].width = min(max_len + 2, 30)

            # 파일 저장
            wb.save(excel_path)
            self.log_message(f"'{keyword}' 결과를 bloggerlevel.xlsx 파일에 저장했습니다.")

        except Exception as e:
            self.log_message(f"엑셀 파일 저장 중 오류 발생: {str(e)}")

    def log_message(self, message):
        """로그 메시지 추가"""
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_text.insert(tk.END, f"[{timestamp}] {message}\n")
        self.log_text.see(tk.END)  # 자동 스크롤
        self.root.update_idletasks()

    def start_processing(self):
        """키워드 처리 시작"""
        keywords_text = self.keyword_entry.get().strip()

        if not keywords_text.strip():
            messagebox.showerror("입력 오류", "키워드를 입력해주세요.\n\n예: 여행, 맛집, IT")
            return

        # 키워드 파싱 (컴마로 분리)
        keywords = [k.strip() for k in keywords_text.split(',') if k.strip()]

        if not keywords:
            messagebox.showerror("입력 오류", "유효한 키워드가 없습니다.")
            return

        # API 키 확인
        if not self.client_id or not self.client_secret:
            messagebox.showerror("API 키 오류", "네이버 API 키(NAVER_CLIENT_ID, NAVER_CLIENT_SECRET)가 설정되지 않았습니다.")
            return

        # 버튼 비활성화
        self.run_button.config(state="disabled")
        self.keyword_entry.config(state="disabled")

        # 백그라운드에서 처리 시작
        thread = threading.Thread(target=self.process_keywords, args=(keywords,))
        thread.daemon = True
        thread.start()

    def process_keywords(self, keywords):
        """키워드들 처리"""
        try:
            self.log_message(f"총 {len(keywords)}개의 키워드를 처리합니다: {', '.join(keywords)}")

            # KST 시간대 설정
            kst = timezone(timedelta(hours=9))
            run_timestamp = datetime.now(kst)

            # 각 키워드 처리
            for idx, keyword in enumerate(keywords, 1):
                self.log_message(f"\n{'='*60}")
                self.log_message(f"[{idx}/{len(keywords)}] '{keyword}' 키워드 처리 시작...")
                self.log_message(f"{'='*60}")

                try:
                    # 블로거 데이터 수집
                    self.log_message(f"'{keyword}'에 대한 블로거 정보를 수집합니다...")
                    blogger_data = get_blogger_data_from_main_page(
                        keyword=keyword,
                        client_id=self.client_id,
                        client_secret=self.client_secret,
                        run_timestamp=run_timestamp,
                        top_n=5,
                        debug=False,
                        debug_file=None,
                    )

                    # 결과 출력
                    self.log_message(f"\n=== '{keyword}' 검색 결과 ===")
                    self.log_message(f"실행 시각: {run_timestamp.strftime('%Y-%m-%d %H:%M:%S')} (KST)")

                    for i, blogger_info in enumerate(blogger_data, 1):
                        blogger = blogger_info["blogger"]
                        today = blogger_info["today"]
                        est = blogger_info["est"]

                        self.log_message(f"{i}위 블로거: {blogger or 'N/A'}")
                        self.log_message(f"    오늘 조회수: {today if today is not None else 'N/A'}")
                        self.log_message(f"    예상 일일 조회수: {est if est is not None else 'N/A'}")

                    # 엑셀 파일에 저장
                    self.log_message(f"'{keyword}' 결과를 엑셀 파일에 저장합니다...")
                    try:
                        self.save_blogger_results_to_excel(keyword, run_timestamp, blogger_data)
                        self.log_message(f"'{keyword}' 엑셀 저장 완료!")
                    except Exception as excel_error:
                        self.log_message(f"경고: '{keyword}' 엑셀 저장 실패 - {str(excel_error)}")
                        self.log_message("프로그램은 계속 실행됩니다.")

                    # 다음 키워드 전 대기
                    if idx < len(keywords):
                        self.log_message("다음 키워드 처리 전 2초 대기...")
                        import time
                        time.sleep(2)

                except Exception as e:
                    self.log_message(f"키워드 '{keyword}' 처리 중 오류 발생: {str(e)}")
                    continue

            self.log_message("\n모든 키워드 처리가 완료되었습니다!")

        except Exception as e:
            self.log_message(f"처리 중 오류 발생: {str(e)}")

        finally:
            # 버튼 재활성화
            self.root.after(0, lambda: self.run_button.config(state="normal"))
            self.root.after(0, lambda: self.keyword_entry.config(state="normal"))

def main():
    root = tk.Tk()
    app = BloggerRankGUI(root)
    root.mainloop()

if __name__ == "__main__":
    main()