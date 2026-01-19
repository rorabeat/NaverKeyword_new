# -*- coding: utf-8 -*-
"""Naver Right Side Related Keywords - Multi Level Extraction

이 스크립트는 네이버 우측 연관 검색어를 sub level3 단계까지 추출합니다.
- Level 1: 초기 seed 키워드들
- Level 2: Level 1의 우측 연관 검색어들
- Level 3: Level 2의 우측 연관 검색어들

각 단계마다 rightside_result 시트에 결과를 누적 저장합니다.
"""

import re
import time
import random
import logging
from datetime import datetime, timezone, timedelta
from urllib.parse import quote_plus
from typing import List, Set, Optional

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


# 설정값들
RESULT_DIR = "result"
LOG_DIR = "result/logs"


def safe_filename(text: str, max_len: int = 40) -> str:
    """파일명으로 안전한 문자열로 변환"""
    text = re.sub(r"[^0-9a-zA-Z가-힣]+", "_", text).strip("_")
    return (text[:max_len] if text else "keyword")


def fetch_naver_search_html(query: str) -> str:
    """
    네이버 통합검색 페이지 HTML 가져오기
    """
    url = f"https://search.naver.com/search.naver?query={quote_plus(query)}"
    headers = {
        # 브라우저처럼 보이게(차단 확률 감소)
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        ),
        "Accept-Language": "ko-KR,ko;q=0.9,en;q=0.8",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Referer": "https://search.naver.com/",
    }

    r = requests.get(url, headers=headers, timeout=15)
    r.raise_for_status()
    return r.text


def extract_related_keywords_from_html(html: str) -> List[str]:
    """
    우측 연관 검색어 추출
    <div class="related_srch">
      ...
      <div class="tit">오키나와 렌트카 비용</div>
    를 기준으로 텍스트를 뽑는다.
    """
    soup = BeautifulSoup(html, "lxml")

    # 1) 가장 정확: related_srch 영역 안의 tit
    nodes = soup.select("div.related_srch div.tit")
    keywords = [n.get_text(strip=True) for n in nodes if n.get_text(strip=True)]

    # 2) 방어 로직: 혹시 클래스가 바뀌면, related_srch 내 keyword 링크 텍스트도 시도
    if not keywords:
        nodes2 = soup.select("div.related_srch a.keyword")
        keywords = [n.get_text(strip=True) for n in nodes2 if n.get_text(strip=True)]

    # 중복 제거(순서 유지)
    deduped = []
    seen = set()
    for k in keywords:
        if k not in seen:
            seen.add(k)
            deduped.append(k)

    return deduped


def setup_logging() -> logging.Logger:
    """로깅 설정"""
    import os
    os.makedirs(LOG_DIR, exist_ok=True)
    os.makedirs(RESULT_DIR, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_filename = f"{LOG_DIR}/rightside_{timestamp}.log"

    logger = logging.getLogger('rightside')
    logger.setLevel(logging.DEBUG)

    # 파일 핸들러
    fh = logging.FileHandler(log_filename, encoding='utf-8')
    fh.setLevel(logging.DEBUG)

    # 콘솔 핸들러
    ch = logging.StreamHandler()
    ch.setLevel(logging.INFO)

    # 포맷터
    formatter = logging.Formatter('%(asctime)s | %(levelname)s | %(message)s')
    fh.setFormatter(formatter)
    ch.setFormatter(formatter)

    logger.addHandler(fh)
    logger.addHandler(ch)

    return logger


def save_to_excel_sheet(seed_keyword: str, related_keywords: List[str], level: int, logger: logging.Logger, is_first_save: bool = False) -> str:
    """keywordList_all 엑셀파일의 rightside_results 시트에 결과를 누적해서 저장"""
    filename = f"{RESULT_DIR}/keywordList_all.xlsx"

    if is_first_save:
        # 첫 번째 저장: 기존 파일이 있으면 로드, 없으면 새로 생성
        try:
            wb = load_workbook(filename)
            if "rightside_results" not in wb.sheetnames:
                # 시트가 없으면 새로 생성
                ws = wb.create_sheet("rightside_results", 0)  # 첫 번째 시트로 생성
                # 헤더 추가
                headers = ["seed", "related_keyword", "level"]
                ws.append(headers)
                # 헤더 스타일
                header_font = Font(bold=True)
                for cell in ws[1]:
                    cell.font = header_font
                    cell.alignment = Alignment(vertical="center")
                # 열 너비 설정
                ws.column_dimensions["A"].width = 30
                ws.column_dimensions["B"].width = 50
                ws.column_dimensions["C"].width = 10
            else:
                # 시트가 있으면 해당 시트 사용
                ws = wb["rightside_results"]
        except FileNotFoundError:
            # 파일이 없으면 새로 생성
            wb = Workbook()
            ws = wb.create_sheet("rightside_results", 0)  # 첫 번째 시트로 생성
            # 헤더
            headers = ["seed", "related_keyword", "level"]
            ws.append(headers)
            # 헤더 스타일
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")
            # 열 너비 설정
            ws.column_dimensions["A"].width = 30
            ws.column_dimensions["B"].width = 50
            ws.column_dimensions["C"].width = 10
    else:
        # 기존 파일 열기
        try:
            wb = load_workbook(filename)
            if "rightside_results" not in wb.sheetnames:
                # 시트가 없으면 새로 생성
                ws = wb.create_sheet("rightside_results")
                # 헤더 추가
                headers = ["seed", "related_keyword", "level"]
                ws.append(headers)
                header_font = Font(bold=True)
                for cell in ws[1]:
                    cell.font = header_font
                    cell.alignment = Alignment(vertical="center")
                ws.column_dimensions["A"].width = 30
                ws.column_dimensions["B"].width = 50
                ws.column_dimensions["C"].width = 10
            else:
                # 시트가 있으면 해당 시트 사용
                ws = wb["rightside_results"]
        except FileNotFoundError:
            # 파일이 없으면 새로 생성
            wb = Workbook()
            ws = wb.create_sheet("rightside_results", 0)

            headers = ["seed", "related_keyword", "level"]
            ws.append(headers)
            header_font = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_font
                cell.alignment = Alignment(vertical="center")
            ws.column_dimensions["A"].width = 30
            ws.column_dimensions["B"].width = 50
            ws.column_dimensions["C"].width = 10

    # 데이터 추가 (마지막 행 다음에)
    for keyword in related_keywords:
        ws.append([seed_keyword, keyword, level])

    # 필터링 설정 업데이트
    max_row = ws.max_row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{max_row}"

    # 파일 저장
    try:
        wb.save(filename)
    except PermissionError as e:
        error_msg = f"엑셀 파일이 열려있어서 저장할 수 없습니다. 파일을 닫고 다시 시도해주세요: {filename}"
        print(f"❌ {error_msg}")
        if logger:
            logger.error(f"파일 저장 실패 - PermissionError: {error_msg}")
        raise RuntimeError(error_msg) from e
    except Exception as e:
        error_msg = f"엑셀 파일 저장 중 오류 발생: {e}"
        print(f"❌ {error_msg}")
        if logger:
            logger.error(f"파일 저장 실패: {error_msg}")
        raise

    if logger:
        logger.info(f"rightside_results 시트 누적 저장: {filename} - seed '{seed_keyword}'의 연관검색어 {len(related_keywords)}개 추가 (레벨 {level})")

    return filename


def extract_multilevel_rightside_related(seeds: List[str], max_level: int = 3, logger: Optional[logging.Logger] = None) -> None:
    """다단계 우측 연관 검색어 추출"""
    try:
        if logger:
            logger.info("=== 다단계 우측 연관 검색어 추출 시작 ===")
            logger.info(f"초기 seed 개수: {len(seeds)}")
            logger.info(f"최대 레벨: {max_level}")

        # 레벨별 키워드 저장
        level_keywords = {1: set(seeds)}  # 중복 방지용 set 사용
        all_processed = set()  # 이미 처리한 키워드들
        is_first_save = True  # 첫 번째 저장인지 확인

        for level in range(1, max_level + 1):
            if logger:
                logger.info(f"=== 레벨 {level} 처리 시작 ===")

            current_level_seeds = list(level_keywords.get(level, []))

            for seed_idx, seed in enumerate(current_level_seeds, 1):
                if seed in all_processed:
                    if logger:
                        logger.debug(f"이미 처리된 seed 건너뜀: {seed}")
                    continue

                if logger:
                    logger.info(f"[레벨 {level}] seed {seed_idx}/{len(current_level_seeds)}: {seed}")

                try:
                    # 우측 연관 검색어 추출
                    html = fetch_naver_search_html(seed)
                    related_keywords = extract_related_keywords_from_html(html)

                    if related_keywords:
                        # 실시간으로 엑셀 시트에 누적 저장
                        excel_file = save_to_excel_sheet(seed, related_keywords, level, logger, is_first_save)
                        is_first_save = False  # 첫 번째 저장 이후에는 False

                        # 로그에 기록
                        if logger:
                            logger.info(f"seed '{seed}' 완료 - 우측 연관 검색어 {len(related_keywords)}개 추출 및 저장 (레벨 {level})")
                            logger.info(f"현재 엑셀 파일: {excel_file}")

                        # 다음 레벨을 위한 키워드 수집 (중복 제거)
                        if level < max_level:
                            next_level_keywords = level_keywords.get(level + 1, set())
                            next_level_keywords.update(related_keywords)
                            level_keywords[level + 1] = next_level_keywords
                    else:
                        if logger:
                            logger.warning(f"seed '{seed}' - 우측 연관 검색어 없음")

                except requests.HTTPError as e:
                    if logger:
                        logger.error(f"HTTP 오류 - seed '{seed}': {e}")
                except Exception as e:
                    if logger:
                        logger.error(f"처리 오류 - seed '{seed}': {e}")

                # 처리 완료된 seed 기록
                all_processed.add(seed)

                # 잠시 대기 (서버 부하 방지)
                time.sleep(random.uniform(1.0, 2.0))

        if logger:
            logger.info("=== 다단계 우측 연관 검색어 추출 완료 ===")

    except Exception as e:
        if logger:
            logger.error(f"프로그램 실행 중 오류 발생: {e}")
        raise


def main():
    """메인 함수"""
    logger = setup_logging()

    try:
        logger.info("=== 네이버 우측 연관 검색어 다단계 추출기 시작 ===")

        # 시드 키워드 입력
        print("시드 키워드를 입력하세요. 여러 개면 콤마(,)로 구분하거나 줄바꿈으로 입력")
        print("입력 끝내려면 빈 줄에서 Enter\n")

        seeds: List[str] = []
        while True:
            line = input("seed > ").strip()
            if not line:
                break
            parts = [p.strip() for p in line.split(",") if p.strip()]
            seeds.extend(parts)

        # 중복 제거하면서 순서 유지
        seen = set()
        seeds = [x for x in seeds if not (x in seen or seen.add(x))]

        if not seeds:
            print("입력된 시드 키워드가 없습니다.")
            return

        print(f"\n처리할 시드 키워드들: {seeds}")
        logger.info(f"입력된 시드 키워드: {seeds}")

        # 다단계 추출 실행
        extract_multilevel_rightside_related(seeds, max_level=3, logger=logger)

    except KeyboardInterrupt:
        logger.info("사용자에 의해 중단됨")
    except Exception as e:
        logger.error(f"프로그램 실행 중 오류 발생: {e}")
        raise


if __name__ == "__main__":
    main()