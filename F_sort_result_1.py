# -*- coding: utf-8 -*-
"""
최근 30일 데이터 정렬 및 분석
- recent30days 시트를 recent30days_sorted 시트에 복사
- 월간 블로그 발행 포화도 계산 및 정렬
- 조건에 맞는 행 강조 표시
"""

import os
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def process_recent30days_data():
    """recent30days 데이터를 정렬하고 분석"""

    filename = "result/keywordList_all.xlsx"

    if not os.path.exists(filename):
        print(f"[ERROR] 파일이 존재하지 않습니다: {filename}")
        return

    try:
        # 기존 파일 로드
        wb = load_workbook(filename)
        print(f"[INFO] 파일 로드 완료: {filename}")

        # recent30days 시트 확인
        if "recent30days" not in wb.sheetnames:
            print("[ERROR] recent30days 시트가 존재하지 않습니다.")
            return

        # 기존 recent30days_sorted 시트가 있으면 삭제
        if "recent30days_sorted" in wb.sheetnames:
            wb.remove(wb["recent30days_sorted"])
            print("[INFO] 기존 recent30days_sorted 시트 삭제")

        # recent30days 시트 복사
        ws_source = wb["recent30days"]
        ws_sorted = wb.copy_worksheet(ws_source)
        ws_sorted.title = "recent30days_sorted"
        print("[INFO] recent30days 시트를 recent30days_sorted로 복사 완료")

        # 데이터 행 수 계산 (헤더 제외)
        max_row = ws_sorted.max_row
        print(f"[INFO] 데이터 행 수: {max_row - 1}개")

        # E열 헤더 추가 (월간 블로그 발행 포화도)
        ws_sorted.cell(row=1, column=5).value = "월간 블로그 발행 포화도"

        # F열 헤더 추가 (추출된 키워드)
        ws_sorted.cell(row=1, column=6).value = "추출된 키워드"

        # 각 행에 대해 계산 수행
        for row in range(2, max_row + 1):  # 헤더 제외
            try:
                # C열과 D열 값 읽기
                c_value = ws_sorted.cell(row=row, column=3).value  # C열
                d_value = ws_sorted.cell(row=row, column=4).value  # D열

                # 숫자 변환 (문자열인 경우 처리)
                try:
                    c_num = float(c_value) if c_value is not None else 0

                    # D열 값 변환: "100+" 같은 문자열 처리
                    if d_value == "100+":
                        d_num = 100.0
                    else:
                        d_num = float(d_value) if d_value is not None else 0
                except (ValueError, TypeError):
                    c_num = 0
                    d_num = 0

                # E열 계산: D/C * 100, 소수점 한 자리
                if c_num != 0:
                    saturation = round((d_num / c_num) * 100, 1)
                else:
                    saturation = 0.0

                ws_sorted.cell(row=row, column=5).value = saturation

                # 추출된 키워드 행 강조 (F열에 "O"가 있는 행만 노란색으로 칠하기)
                if c_num > 500 and d_num < 100 and saturation < 10:
                    # F열에 "O" 표시
                    ws_sorted.cell(row=row, column=6).value = "O"

                # F열에 "O"가 있는 행만 노란색으로 강조
                if ws_sorted.cell(row=row, column=6).value == "O":
                    # 노란색 배경 적용
                    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

                    # 해당 행의 모든 셀에 노란색 적용 (A열부터 E열까지)
                    for col in range(1, 6):  # A부터 E열까지
                        cell = ws_sorted.cell(row=row, column=col)
                        cell.fill = yellow_fill

            except Exception as e:
                print(f"[WARN] {row}행 처리 중 오류: {e}")
                continue

        # E열 기준 오름차순 정렬 (헤더 제외)
        print("[INFO] E열 기준 오름차순 정렬 중...")

        # 데이터 행만 추출하여 정렬 (원본 데이터와 함께 강조 정보도 저장)
        data_rows = []
        for row in range(2, max_row + 1):
            row_data = []
            for col in range(1, 7):  # A부터 F열까지
                row_data.append(ws_sorted.cell(row=row, column=col).value)

            # 강조 표시 여부도 저장 (F열에 "O"가 있는지)
            is_highlighted = row_data[5] == "O"  # F열 값 (인덱스 5 = 6번째 열)
            data_rows.append((row_data, is_highlighted))

        # E열(인덱스 4) 기준으로 정렬
        data_rows.sort(key=lambda x: x[0][4] if x[0][4] is not None else 0)

        # 정렬된 데이터를 다시 시트에 쓰기
        for row_idx, (row_data, is_highlighted) in enumerate(data_rows, start=2):
            for col_idx, value in enumerate(row_data, start=1):
                ws_sorted.cell(row=row_idx, column=col_idx).value = value

            # 강조 표시 적용 (전체 행에 대해 한 번에 처리)
            if is_highlighted:
                yellow_fill = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
                for col in range(1, 6):  # A부터 E열까지
                    ws_sorted.cell(row=row_idx, column=col).fill = yellow_fill


        # 열 너비 및 스타일 설정
        ws_sorted.column_dimensions["A"].width = 20
        ws_sorted.column_dimensions["B"].width = 30
        ws_sorted.column_dimensions["C"].width = 15
        ws_sorted.column_dimensions["D"].width = 15
        ws_sorted.column_dimensions["E"].width = 25
        ws_sorted.column_dimensions["F"].width = 15

        # 헤더 스타일링
        header_font = Font(bold=True)
        center_alignment = Alignment(horizontal="center")

        for col in range(1, 7):  # A부터 F열까지
            cell = ws_sorted.cell(row=1, column=col)
            cell.font = header_font
            cell.alignment = center_alignment

        # 필터링 설정
        ws_sorted.freeze_panes = "A2"
        ws_sorted.auto_filter.ref = f"A1:F{max_row}"

        # 파일 저장
        wb.save(filename)

        print("[SUCCESS] 데이터 처리 완료!")
        print(f"   - 원본 시트: recent30days")
        print(f"   - 새 시트: recent30days_sorted")
        print(f"   - 저장 파일: {filename}")
        print(f"   - 처리된 행 수: {max_row - 1}개")

    except Exception as e:
        print(f"[ERROR] 오류 발생: {e}")


def main():
    """메인 함수"""
    print("=== 최근 30일 데이터 정렬 및 분석 시작 ===")
    process_recent30days_data()
    print("=== 최근 30일 데이터 정렬 및 분석 완료 ===")


if __name__ == "__main__":
    main()