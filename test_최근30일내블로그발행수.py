import os
import re
import requests
from datetime import datetime, timedelta, timezone
from typing import Optional, Tuple, List, Dict, Any

from dotenv import load_dotenv  # pip install python-dotenv

# Excel 출력용
# pip install openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

NAVER_BLOG_URL = "https://openapi.naver.com/v1/search/blog.json"


def fetch_naver_blog_json(
    query: str,
    client_id: str,
    client_secret: str,
    display: int = 100,
    start: int = 1,
    sort: str = "date",
) -> Dict[str, Any]:
    headers = {
        "X-Naver-Client-Id": client_id,
        "X-Naver-Client-Secret": client_secret,
    }
    params = {"query": query, "display": display, "start": start, "sort": sort}

    r = requests.get(NAVER_BLOG_URL, headers=headers, params=params, timeout=10)
    r.raise_for_status()
    return r.json()


def strip_html(text: str) -> str:
    if not text:
        return ""
    # title에 <b> 같은 태그가 포함될 수 있어 간단 제거
    return re.sub(r"<[^>]+>", "", text)


def recent_30d_count_capped(
    query: str,
    client_id: str,
    client_secret: str,
    first_data: Optional[Dict[str, Any]] = None,
    limit: int = 100,
    days: int = 30,
) -> Tuple[int, bool, int]:
    """최근 30일 발행 글 수를 세되,
    - 30일 이전 글이 나오면 중단
    - limit(기본 100) 도달하면 즉시 중단하고 over_100=True

    return: (count, over_100, cutoff_yyyymmdd)
    """

    kst = timezone(timedelta(hours=9))
    cutoff = (datetime.now(kst) - timedelta(days=days)).date()
    cutoff_yyyymmdd = int(cutoff.strftime("%Y%m%d"))

    count = 0
    start = 1
    display = 100

    def consume_items(items: List[Dict[str, Any]]) -> Tuple[bool, bool]:
        """return (stop_by_date, reached_limit)"""
        nonlocal count

        for item in items:
            postdate_str = item.get("postdate")  # "YYYYMMDD"
            if not postdate_str or len(str(postdate_str)) != 8:
                continue

            if int(postdate_str) < cutoff_yyyymmdd:
                return True, False

            count += 1
            if count >= limit:
                return True, True

        return False, False

    # 1) 첫 페이지 데이터가 이미 있으면 그걸 먼저 소비
    if first_data is not None:
        items = first_data.get("items", []) or []
        stop_by_date, reached_limit = consume_items(items)
        if reached_limit:
            return limit, True, cutoff_yyyymmdd
        if stop_by_date:
            return count, False, cutoff_yyyymmdd
        start = 101  # 첫 페이지를 처리했으니 다음 페이지

    # 2) 필요하면 다음 페이지들 조회
    while True:
        data = fetch_naver_blog_json(
            query=query,
            client_id=client_id,
            client_secret=client_secret,
            display=display,
            start=start,
            sort="date",
        )
        items = data.get("items", []) or []
        if not items:
            break

        stop_by_date, reached_limit = consume_items(items)
        if reached_limit:
            return limit, True, cutoff_yyyymmdd
        if stop_by_date:
            break

        start += display
        if start > 1000:  # 안전장치(원하면 조정/삭제)
            break

    return count, False, cutoff_yyyymmdd


def write_excel_report(
    summary_rows: List[Dict[str, Any]],
    item_rows: List[Dict[str, Any]],
    out_path: str,
) -> None:
    wb = Workbook()

    # -------------------- Summary sheet --------------------
    ws = wb.active
    ws.title = "summary"

    headers = [
        "keyword",
        "cutoff_yyyymmdd",
        "blog_last30_display",
        "blog_last30_count",
        "blog_over_100",
        "api_total",
        "api_lastBuildDate",
        "error",
    ]
    ws.append(headers)

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for r in summary_rows:
        ws.append([
            r.get("keyword", ""),
            r.get("cutoff_yyyymmdd", ""),
            r.get("blog_last30_display", ""),
            r.get("blog_last30_count", ""),
            r.get("blog_over_100", ""),
            r.get("api_total", ""),
            r.get("api_lastBuildDate", ""),
            r.get("error", ""),
        ])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{chr(ord('A') + len(headers) - 1)}{len(summary_rows) + 1}"

    # 컬럼 너비 자동(너무 길면 제한)
    for col_idx, col_name in enumerate(headers, start=1):
        max_len = len(col_name)
        for row_idx in range(2, len(summary_rows) + 2):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws.column_dimensions[chr(64 + col_idx)].width = min(max_len + 2, 60)

    # -------------------- Items sheet (상위 N개 샘플) --------------------
    ws2 = wb.create_sheet("items_sample")

    headers2 = [
        "keyword",
        "rank",
        "postdate",
        "bloggername",
        "title",
        "link",
    ]
    ws2.append(headers2)
    for cell in ws2[1]:
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for r in item_rows:
        ws2.append([
            r.get("keyword", ""),
            r.get("rank", ""),
            r.get("postdate", ""),
            r.get("bloggername", ""),
            r.get("title", ""),
            r.get("link", ""),
        ])

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{chr(ord('A') + len(headers2) - 1)}{len(item_rows) + 1}"

    # items 시트 너비
    for col_idx, col_name in enumerate(headers2, start=1):
        max_len = len(col_name)
        for row_idx in range(2, len(item_rows) + 2):
            v = ws2.cell(row=row_idx, column=col_idx).value
            if v is None:
                continue
            max_len = max(max_len, len(str(v)))
        ws2.column_dimensions[chr(64 + col_idx)].width = min(max_len + 2, 80)

    wb.save(out_path)


def read_keywords_from_user() -> List[str]:
    print("키워드를 입력하세요. (여러 개면 콤마(,)로 구분하거나 줄바꿈으로 입력)")
    print("입력 끝내려면 빈 줄에서 Enter를 누르세요.\n")

    keywords: List[str] = []
    while True:
        line = input("키워드 입력 > ").strip()
        if not line:
            break

        parts = [p.strip() for p in line.split(",") if p.strip()]
        keywords.extend(parts)

    # 중복 제거(입력 순서 유지)
    deduped: List[str] = []
    seen = set()
    for k in keywords:
        if k not in seen:
            seen.add(k)
            deduped.append(k)

    return deduped


def main() -> None:
    load_dotenv()

    client_id = os.getenv("NAVER_CLIENT_ID", "").strip()
    client_secret = os.getenv("NAVER_CLIENT_SECRET", "").strip()

    if not client_id or not client_secret:
        print("ERROR: .env 또는 환경변수에 NAVER_CLIENT_ID / NAVER_CLIENT_SECRET 이 없습니다.")
        print("프로젝트 폴더에 .env 파일을 만들고 다음을 넣으세요:")
        print("NAVER_CLIENT_ID=...")
        print("NAVER_CLIENT_SECRET=...")
        return

    keywords = read_keywords_from_user()
    if not keywords:
        print("키워드가 입력되지 않았습니다. 종료합니다.")
        return

    SHOW_ITEMS_N = 5  # items_sample 시트에 각 키워드당 샘플로 저장할 개수

    kst = timezone(timedelta(hours=9))
    ts = datetime.now(kst).strftime("%Y%m%d_%H%M%S")
    out_path = f"naver_blog_last30_{ts}.xlsx"

    summary_rows: List[Dict[str, Any]] = []
    item_rows: List[Dict[str, Any]] = []

    print(f"실행시각(KST): {datetime.now(kst).strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"키워드 개수: {len(keywords)}")

    for kw in keywords:
        row: Dict[str, Any] = {
            "keyword": kw,
            "cutoff_yyyymmdd": "",
            "blog_last30_display": "",
            "blog_last30_count": "",
            "blog_over_100": "",
            "api_total": "",
            "api_lastBuildDate": "",
            "error": "",
        }

        try:
            # 1) 첫 페이지 조회(최신순)
            data1 = fetch_naver_blog_json(
                query=kw,
                client_id=client_id,
                client_secret=client_secret,
                display=100,
                start=1,
                sort="date",
            )

            # 2) 최근 30일 카운트(100에서 캡, 100+ 여부 플래그)
            count_30d, over_100, cutoff_yyyymmdd = recent_30d_count_capped(
                query=kw,
                client_id=client_id,
                client_secret=client_secret,
                first_data=data1,
                limit=100,
                days=30,
            )

            row["cutoff_yyyymmdd"] = cutoff_yyyymmdd
            row["blog_last30_count"] = count_30d
            row["blog_over_100"] = bool(over_100)
            row["blog_last30_display"] = "100+" if over_100 else str(count_30d)
            row["api_total"] = data1.get("total", "")
            row["api_lastBuildDate"] = data1.get("lastBuildDate", "")

            # 3) 샘플 items 저장(첫 페이지 상위 N개)
            items = data1.get("items", []) or []
            for i, item in enumerate(items[:SHOW_ITEMS_N], start=1):
                item_rows.append({
                    "keyword": kw,
                    "rank": i,
                    "postdate": item.get("postdate", ""),
                    "bloggername": item.get("bloggername", ""),
                    "title": strip_html(item.get("title", "")),
                    "link": item.get("link", ""),
                })

            print(f"- {kw}: {row['blog_last30_display']} (cutoff {cutoff_yyyymmdd})")

        except requests.HTTPError as e:
            body = ""
            try:
                body = e.response.text[:500]
            except Exception:
                pass
            row["error"] = f"HTTPError: {e} | body: {body}"
            print(f"- {kw}: HTTPError")

        except Exception as e:
            row["error"] = f"Error: {repr(e)}"
            print(f"- {kw}: Error")

        summary_rows.append(row)

    write_excel_report(summary_rows, item_rows, out_path)
    print(f"\n완료! 결과를 '{out_path}' 엑셀 파일로 저장했습니다.")


if __name__ == "__main__":
    main()
