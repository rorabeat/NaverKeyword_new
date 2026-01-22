import openpyxl
wb = openpyxl.load_workbook('result/keywordList_all.xlsx')
ws = wb['recent30days_sorted']
# 마지막 행의 F열을 'O'로 설정
last_row = ws.max_row
ws.cell(row=last_row, column=6).value = 'O'
wb.save('result/keywordList_all.xlsx')
keyword = ws.cell(row=last_row, column=2).value
print(f'행 {last_row}의 F열을 "O"로 설정했습니다. 키워드: {keyword}')