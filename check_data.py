from openpyxl import load_workbook

wb = load_workbook('result/keywordList_all.xlsx')
ws = wb['recent30days']
print('recent30days 시트 샘플 데이터:')
for row in range(2, 8):
    c_val = ws.cell(row=row, column=3).value
    d_val = ws.cell(row=row, column=4).value
    print('행{}: C={} (type: {}), D={} (type: {})'.format(
        row, c_val, type(c_val).__name__, d_val, type(d_val).__name__))

# C열이 500보다 큰 행 찾기
print('\nC열이 500보다 큰 행들:')
large_c_rows = []
for row in range(2, ws.max_row + 1):
    c_val = ws.cell(row=row, column=3).value
    d_val = ws.cell(row=row, column=4).value
    try:
        c_num = float(c_val) if c_val is not None else 0
        if c_num > 500:
            large_c_rows.append((row, c_val, d_val))

    except (ValueError, TypeError):
        continue

if large_c_rows:
    for row, c_val, d_val in large_c_rows[:5]:  # 처음 5개만
        print('행{}: C={}, D={}'.format(row, c_val, d_val))
    if len(large_c_rows) > 5:
        print('... 외 {}개'.format(len(large_c_rows) - 5))
else:
    print('없음')