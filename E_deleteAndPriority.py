# -*- coding: utf-8 -*-
"""searchResult 시트에서 relKeyword 중복 제거 및 데이터 추출

searchResult 시트의 relKeyword 열(B열)에서 중복을 제거하고,
A열(seed_keyword), B열(relKeyword), E열(monthlyTotal) 정보만 수집하여
removeDuplicate 시트에 저장합니다.
"""

import os
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


def remove_duplicate_relkeywords():
    """searchResult 시트에서 중복 relKeyword 제거 및 데이터 추출"""

    filename = "result/keywordList_all.xlsx"

    if not os.path.exists(filename):
        print(f"[ERROR] 파일이 존재하지 않습니다: {filename}")
        return

    try:
        # 기존 파일 로드
        wb = load_workbook(filename)
        print(f"[INFO] 파일 로드 완료: {filename}")

        # searchResult 시트 확인
        if "searchResult" not in wb.sheetnames:
            print("[ERROR] searchResult 시트가 존재하지 않습니다. 먼저 검색결과 수집을 실행해주세요.")
            return

        ws_search = wb["searchResult"]
        print(f"[INFO] searchResult 시트에서 데이터 읽기 중... ({ws_search.max_row}행)")

        # 데이터 수집 및 중복 제거
        unique_keywords = {}  # relKeyword를 키로 하는 딕셔너리

        # 헤더를 제외한 데이터 행부터 처리
        for row in range(2, ws_search.max_row + 1):
            seed_keyword = ws_search.cell(row=row, column=1).value  # A열
            rel_keyword = ws_search.cell(row=row, column=2).value   # B열
            monthly_total = ws_search.cell(row=row, column=5).value # E열

            # 빈 값이나 None은 건너뛰기
            if not rel_keyword:
                continue

            # relKeyword를 키로 하여 중복 제거 (첫 번째 발견된 값 유지)
            if rel_keyword not in unique_keywords:
                unique_keywords[rel_keyword] = {
                    'seed_keyword': seed_keyword,
                    'relKeyword': rel_keyword,
                    'monthlyTotal': monthly_total
                }

        if not unique_keywords:
            print("[ERROR] searchResult 시트에 처리할 데이터가 없습니다.")
            return

        print(f"[INFO] 중복 제거 전 총 {ws_search.max_row - 1}개 → 중복 제거 후 {len(unique_keywords)}개 키워드")

        # removeDuplicate 시트 생성 (기존 시트가 있으면 제거)
        if "removeDuplicate" in wb.sheetnames:
            wb.remove(wb["removeDuplicate"])
            print("[INFO] 기존 removeDuplicate 시트 제거")

        ws_remove_dup = wb.create_sheet("removeDuplicate")
        print("[INFO] removeDuplicate 시트 생성")

        # 헤더 추가
        headers = ["seed_keyword", "relKeyword", "monthlyTotal"]
        ws_remove_dup.append(headers)

        # 헤더 스타일링
        header_font = Font(bold=True)
        for cell in ws_remove_dup[1]:
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")

        # 열 너비 설정
        ws_remove_dup.column_dimensions["A"].width = 25
        ws_remove_dup.column_dimensions["B"].width = 40
        ws_remove_dup.column_dimensions["C"].width = 15

        # 중복 제거된 데이터 추가
        for keyword_data in unique_keywords.values():
            row_data = [
                keyword_data['seed_keyword'],
                keyword_data['relKeyword'],
                keyword_data['monthlyTotal']
            ]
            ws_remove_dup.append(row_data)

        # 필터링 설정
        max_row = ws_remove_dup.max_row
        ws_remove_dup.freeze_panes = "A2"
        ws_remove_dup.auto_filter.ref = f"A1:C{max_row}"

        # 파일 저장
        try:
            wb.save(filename)
            print("[SUCCESS] 중복 제거 및 데이터 추출 완료!")
            print(f"   - 원본 데이터: {ws_search.max_row - 1}개")
            print(f"   - 중복 제거 후: {len(unique_keywords)}개")
            print(f"   - 저장 파일: {filename}")
            print(f"   - 새 시트: removeDuplicate")
        except PermissionError:
            print("[WARN] 엑셀 파일 저장 실패: 파일이 열려있거나 쓰기 권한이 없습니다.")
            print("   - 결과를 확인하려면 엑셀 파일을 닫고 다시 실행해주세요.")
        except Exception as e:
            print(f"[ERROR] 파일 저장 중 오류: {e}")

    except Exception as e:
        print(f"[ERROR] 오류 발생: {e}")


def main():
    """메인 함수"""
    print("=== searchResult 시트 relKeyword 중복 제거 시작 ===")
    remove_duplicate_relkeywords()
    print("=== searchResult 시트 relKeyword 중복 제거 완료 ===")


if __name__ == "__main__":
    main()